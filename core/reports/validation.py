from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook

from core.analysis.schemas import AIAnalysisV1, AnalysisPacketV1
from core.inventory.models import ReportSnapshotV1
from core.reports.models import ReportManifestV1, SHEET_NAMES
from core.reports.view_model import ReportViewModel


class WorkbookValidationError(ValueError):
    pass


EXPECTED_HEADERS = {
    "空间体检概览": ["项目", "内容"],
    "空间明细清单": [
        "对象 ID", "类别", "名称", "位置或父级",
        "对象本身大小", "对象本身大小（原始 byte）",
        "实际占用磁盘", "实际占用磁盘（原始 byte）",
        "可能可腾出", "可能可腾出（原始 byte）",
        "创建线索", "创建来源", "修改线索", "活动线索",
        "活动可信程度", "时间限制", "规则提示", "AI 观察编号",
        "我的决定", "我的备注",
    ],
    "建议优先查看": [
        "优先级", "对象 ID", "对象", "空间影响", "空间影响（原始 byte）",
        "原因", "唯一性风险", "重建难度", "证据可信程度", "转到空间明细清单",
    ],
    "可能重复与安装包": [
        "类型", "组号", "对象 ID", "名称", "位置", "大小",
        "大小（原始 byte）", "版本线索", "依据", "核验状态", "下一步",
    ],
    "报告说明与记录": ["审计字段", "值"],
}


def workbook_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def validate_workbook(
    path: Path,
    *,
    view: ReportViewModel,
    snapshot: ReportSnapshotV1,
    packet: AnalysisPacketV1,
    analysis: AIAnalysisV1,
) -> dict[str, Any]:
    issues: list[str] = []
    with ZipFile(path) as archive:
        bad_member = archive.testzip()
        names = set(archive.namelist())
    if bad_member is not None:
        issues.append(f"ZIP_CORRUPT:{bad_member}")
    if any("externalLinks" in name or "vbaProject" in name for name in names):
        issues.append("EXTERNAL_OR_MACRO_CONTENT")
    workbook = load_workbook(path, data_only=False, read_only=False)
    if workbook.sheetnames != SHEET_NAMES:
        issues.append("SHEET_ORDER")
    for sheet in workbook.worksheets:
        headers = [cell.value for cell in sheet[1]]
        if headers != EXPECTED_HEADERS[sheet.title]:
            issues.append(f"HEADERS:{sheet.title}")
        if sheet.freeze_panes != "A2":
            issues.append(f"FREEZE:{sheet.title}")
        if not sheet.auto_filter.ref:
            issues.append(f"FILTER:{sheet.title}")
        if sheet.merged_cells.ranges:
            issues.append(f"MERGED_DATA:{sheet.title}")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    issues.append(f"FORMULA_CELL:{sheet.title}:{cell.coordinate}")
    overview = workbook[SHEET_NAMES[0]]
    if overview["B4"].is_date is not True:
        issues.append("GENERATED_AT_NOT_DATE")
    audit = workbook[SHEET_NAMES[4]]
    audit_values = {
        str(audit.cell(row, 1).value): audit.cell(row, 2).value
        for row in range(2, audit.max_row + 1)
    }
    expected_ids = {
        "report_id": view.report_id,
        "artifact_id": view.artifact_id,
        "snapshot_id": snapshot.snapshot_id,
        "packet_id": packet.packet_id,
        "analysis_id": analysis.analysis_id,
    }
    for key, expected in expected_ids.items():
        if audit_values.get(key) != expected:
            issues.append(f"ID_MISMATCH:{key}")
    priority = workbook[SHEET_NAMES[2]]
    for row in range(2, priority.max_row + 1):
        link = priority.cell(row, 10).hyperlink
        if link is None or not str(link.target).startswith(f"#'{SHEET_NAMES[1]}'!A"):
            issues.append(f"INTERNAL_LINK:{row}")
    inventory = workbook[SHEET_NAMES[1]]
    for row in range(2, inventory.max_row + 1):
        for column in (6, 8, 10):
            value = inventory.cell(row, column).value
            if value is not None and not isinstance(value, int):
                issues.append(f"RAW_BYTE_TYPE:{row}:{column}")
        for column in (11, 13, 14):
            value = inventory.cell(row, column).value
            if value != "未知" and not inventory.cell(row, column).is_date:
                issues.append(f"DATE_TYPE:{row}:{column}")
    workbook.close()
    if issues:
        raise WorkbookValidationError(";".join(issues))
    return {
        "passed": True,
        "sheet_names": SHEET_NAMES,
        "sheet_count": len(SHEET_NAMES),
        "zip_integrity": True,
        "reopen_validated": True,
        "formula_errors": 0,
        "id_consistency": True,
        "workbook_sha256": workbook_sha256(path),
    }


def build_report_manifest(
    *,
    path: Path,
    view: ReportViewModel,
    validation: dict[str, Any],
) -> ReportManifestV1:
    return ReportManifestV1(
        report_id=view.report_id,
        artifact_id=view.artifact_id,
        snapshot_id=view.snapshot_id,
        packet_id=view.packet_id,
        analysis_id=view.analysis_id,
        generated_at=view.generated_at,
        artifact_privacy_mode=view.artifact_privacy_mode,
        provider=view.provider,
        analysis_status=view.analysis_status,
        filename=path.name,
        workbook_sha256=validation["workbook_sha256"],
    )
