from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet


HEADER_FILL = PatternFill("solid", fgColor="123B5D")
HEADER_FONT = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
BODY_FONT = Font(name="Microsoft YaHei", color="172331")
AUDIT_FILL = PatternFill("solid", fgColor="E8F1F5")
THIN_BORDER = Border(
    bottom=Side(style="thin", color="C8D5DD"),
)


def style_sheet(worksheet: Worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
    for column_cells in worksheet.columns:
        letter = column_cells[0].column_letter
        width = max(
            10,
            min(
                48,
                max(
                    len(str(cell.value)) if cell.value is not None else 0
                    for cell in column_cells
                )
                + 2,
            ),
        )
        worksheet.column_dimensions[letter].width = width
    worksheet.row_dimensions[1].height = 30
