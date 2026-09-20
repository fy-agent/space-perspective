from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from core.reports.models import SHEET_NAMES
from core.reports.styles import style_sheet
from core.reports.view_model import ReportViewModel


RAW_BYTE_FORMAT = "#,##0"


def friendly_bytes(value: int | None) -> str:
    if value is None:
        return "未评估"
    units = ["B", "KiB", "MiB", "GiB", "TiB"]
    amount = float(value)
    for unit in units:
        if abs(amount) < 1024 or unit == units[-1]:
            return f"{amount:.1f} {unit}" if unit != "B" else f"{int(amount)} B"
        amount /= 1024
    return f"{value} B"


def _append_rows(sheet, headers: list[str], rows: list[list[object]]) -> None:
    sheet.append([_safe_cell_value(value) for value in headers])
    for row in rows:
        sheet.append([_safe_cell_value(value) for value in row])


def _safe_cell_value(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _excel_datetime(value: datetime | None) -> datetime | str:
    return value.replace(tzinfo=None) if value is not None else "未知"


def render_workbook(view: ReportViewModel, output_path: Path) -> Path:
    workbook = Workbook()
    overview = workbook.active
    overview.title = SHEET_NAMES[0]
    generated_at = datetime.fromisoformat(view.generated_at).replace(tzinfo=None)
    _append_rows(
        overview,
        ["项目", "内容"],
        [
            ["报告编号", view.report_id],
            ["制品编号", view.artifact_id],
            ["生成时间", generated_at],
            ["范围", "受控 fixture"],
            ["平台", view.platform],
            ["隐私模式", view.artifact_privacy_mode],
            ["报告状态", view.analysis_status],
            ["对象本身大小", friendly_bytes(view.total_logical_size_bytes)],
            ["对象本身大小（原始 byte）", view.total_logical_size_bytes],
            ["实际占用磁盘", friendly_bytes(view.total_allocated_size_bytes)],
            ["实际占用磁盘（原始 byte）", view.total_allocated_size_bytes],
            ["可能可腾出", friendly_bytes(view.total_reclaimable_estimate_bytes)],
            ["可能可腾出（原始 byte）", view.total_reclaimable_estimate_bytes],
            ["盘点覆盖", view.coverage_text],
            ["权限缺口", view.permission_gap_text],
            ["AI 帮你整理的观察", "仅供判断，不代表已授权处理"],
            ["如何理解这份报告", "空间数字由确定性代码计算；规则、AI 观察和用户决定彼此分列。"],
        ],
    )
    overview["B4"].number_format = "yyyy-mm-dd hh:mm:ss"
    for coordinate in ("B10", "B12", "B14"):
        if type(overview[coordinate].value) is int:
            overview[coordinate].number_format = RAW_BYTE_FORMAT

    inventory = workbook.create_sheet(SHEET_NAMES[1])
    _append_rows(
        inventory,
        [
            "对象 ID", "类别", "名称", "位置或父级",
            "对象本身大小", "对象本身大小（原始 byte）",
            "实际占用磁盘", "实际占用磁盘（原始 byte）",
            "可能可腾出", "可能可腾出（原始 byte）",
            "创建线索", "创建来源", "修改线索", "活动线索",
            "活动可信程度", "时间限制", "规则提示", "AI 观察编号",
            "我的决定", "我的备注",
        ],
        [
            [
                row.object_id,
                row.object_type,
                row.name,
                row.location,
                friendly_bytes(row.logical_size_bytes),
                row.logical_size_bytes,
                friendly_bytes(row.allocated_size_bytes),
                row.allocated_size_bytes,
                friendly_bytes(row.reclaimable_estimate_bytes),
                row.reclaimable_estimate_bytes,
                _excel_datetime(row.created_value),
                row.created_source,
                _excel_datetime(row.modified_value),
                _excel_datetime(row.activity_value),
                row.activity_confidence,
                row.time_limitation,
                row.rule_hint,
                row.ai_observation_ids,
                row.user_decision,
                row.user_note,
            ]
            for row in view.inventory_rows
        ],
    )
    for row_number in range(2, inventory.max_row + 1):
        for column in (6, 8, 10):
            cell = inventory.cell(row_number, column)
            if type(cell.value) is int:
                cell.number_format = RAW_BYTE_FORMAT
        for column in (11, 13, 14):
            cell = inventory.cell(row_number, column)
            if cell.is_date:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"

    priority = workbook.create_sheet(SHEET_NAMES[2])
    _append_rows(
        priority,
        [
            "优先级", "对象 ID", "对象", "空间影响", "空间影响（原始 byte）",
            "原因", "唯一性风险", "重建难度", "证据可信程度", "转到空间明细清单",
        ],
        [
            [
                row.priority,
                row.object_id,
                row.name,
                friendly_bytes(row.logical_size_bytes),
                row.logical_size_bytes,
                row.reason,
                row.uniqueness_risk,
                row.rebuild_difficulty,
                row.evidence_confidence,
                "查看明细",
            ]
            for row in view.priority_rows
        ],
    )
    for excel_row, item in enumerate(view.priority_rows, start=2):
        raw_size_cell = priority.cell(excel_row, 5)
        if type(raw_size_cell.value) is int:
            raw_size_cell.number_format = RAW_BYTE_FORMAT
        cell = priority.cell(excel_row, 10)
        cell.hyperlink = f"#'{SHEET_NAMES[1]}'!A{item.inventory_row}"
        cell.font = Font(name="Microsoft YaHei", color="0563C1", underline="single")

    duplicates = workbook.create_sheet(SHEET_NAMES[3])
    _append_rows(
        duplicates,
        [
            "类型", "组号", "对象 ID", "名称", "位置", "大小",
            "大小（原始 byte）", "版本线索", "依据", "核验状态", "下一步",
        ],
        [
            [
                row.kind,
                row.group_id,
                row.object_id,
                row.name,
                row.location,
                friendly_bytes(row.logical_size_bytes),
                row.logical_size_bytes,
                row.version_hint,
                row.basis,
                row.verification_status,
                row.next_step,
            ]
            for row in view.duplicate_installer_rows
        ],
    )
    for row_number in range(2, duplicates.max_row + 1):
        raw_size_cell = duplicates.cell(row_number, 7)
        if type(raw_size_cell.value) is int:
            raw_size_cell.number_format = RAW_BYTE_FORMAT

    audit = workbook.create_sheet(SHEET_NAMES[4])
    audit_rows = [
        ["report_id", view.report_id],
        ["artifact_id", view.artifact_id],
        ["snapshot_id", view.snapshot_id],
        ["packet_id", view.packet_id],
        ["analysis_id", view.analysis_id],
        ["profile", view.profile],
        ["content_read", False],
        ["hash_mode", "none"],
        ["provider", view.provider],
        ["model", view.model],
        ["analysis_status", view.analysis_status],
        ["analysis_execution_mode", view.analysis_execution_mode],
        ["packet_sha256", view.packet_sha256],
        ["prompt_version", view.prompt_version],
        ["analysis_schema_version", view.analysis_schema_version],
        ["consent_receipt_id", view.consent_receipt_id],
        ["consent_status", view.consent_status],
        ["analysis_budget", view.analysis_budget],
        ["provider_calls", view.provider_calls],
        ["network_calls", view.network_calls],
        ["input_tokens", view.input_tokens],
        ["cached_input_tokens", view.cached_input_tokens],
        ["output_tokens", view.output_tokens],
        ["estimated_cost", view.estimated_cost],
        ["cost_basis", view.cost_basis],
        ["latency_ms", view.latency_ms],
        ["analysis_error_code", view.analysis_error_code],
        ["analysis_fallback", view.analysis_fallback],
        ["privacy_mode", view.artifact_privacy_mode],
        ["scope", "受控 fixture"],
        ["coverage", view.coverage_text],
        ["permission_gaps", view.permission_gap_text],
        ["空间口径", "对象本身大小 / 实际占用磁盘 / 可能可腾出"],
        ["时间证据", "creation 未知时保持未知；atime 仅为低可信活动线索。"],
        ["AI 说明", "AI 帮你整理的观察，仅供判断，不代表已授权处理。"],
        ["文件动作", "未执行"],
        ["网络", "未调用；network_calls=0"],
    ]
    audit_rows.extend(
        [f"限制 {index}", limitation]
        for index, limitation in enumerate(view.limitations, start=1)
    )
    _append_rows(audit, ["审计字段", "值"], audit_rows)

    for sheet in workbook.worksheets:
        style_sheet(sheet)
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False
    workbook.calculation.calcMode = "manual"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
