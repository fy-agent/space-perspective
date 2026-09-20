from __future__ import annotations

import hashlib
from io import BytesIO
import json
from pathlib import Path
import socket
import sqlite3
from unittest.mock import Mock
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from core.app.config import Settings
from core.app.main import create_app
from core.reports.models import SHEET_NAMES
import core.scanner.hashing as p0_hashing
import core.services.inventory as inventory_service_module


ROOT = Path(__file__).parents[2]
FIXTURE_ROOT = ROOT / "fixtures" / "ai-report"


def _fixture_integrity() -> tuple[list[dict[str, object]], int]:
    """Test-only validator; these reads/hashes are never product counters."""
    records: list[dict[str, object]] = []
    hashes = 0
    for path in sorted(
        FIXTURE_ROOT.rglob("*"),
        key=lambda item: item.relative_to(FIXTURE_ROOT).as_posix(),
    ):
        info = path.stat(follow_symlinks=False)
        record: dict[str, object] = {
            "relative_path": path.relative_to(FIXTURE_ROOT).as_posix(),
            "kind": "directory" if path.is_dir() else "file",
            "size": info.st_size,
            "mtime_ns": info.st_mtime_ns,
            "sha256": None,
        }
        if path.is_file():
            digest = hashlib.sha256()
            with path.open("rb") as source:
                while chunk := source.read(1024 * 1024):
                    digest.update(chunk)
            record["sha256"] = digest.hexdigest()
            hashes += 1
        records.append(record)
    return records, hashes


def _create_scan(client: TestClient) -> dict[str, object]:
    response = client.post(
        "/v1/inventory/scans",
        json={
            "fixture_id": "standard",
            "profile": "inventory_metadata",
            "compute_hash": False,
        },
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["status"] == "completed"
    assert payload["fixture_id"] == "standard"
    assert payload["profile"] == "inventory_metadata"
    assert payload["content_read"] is False
    assert payload["hash_mode"] == "none"
    assert payload["safety"] == {
        "product_content_reads": 0,
        "product_file_hashes": 0,
        "network_calls": 0,
        "file_actions": 0,
    }
    return payload


def _create_snapshot(client: TestClient, scan_id: str) -> dict[str, object]:
    response = client.post("/v1/inventory/snapshots", json={"scan_id": scan_id})
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["scan_id"] == scan_id
    assert payload["content_read"] is False
    assert payload["hash_mode"] == "none"
    return payload


def _preview_packet(
    client: TestClient,
    snapshot_id: str,
    *,
    privacy: str,
    execution: str,
) -> dict[str, object]:
    response = client.post(
        "/v1/inventory/analysis-packets/preview",
        json={
            "snapshot_id": snapshot_id,
            "artifact_privacy_mode": privacy,
            "analysis_execution_mode": execution,
            "top_k": 200,
        },
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["packet"]["snapshot_id"] == snapshot_id
    assert payload["packet"]["artifact_privacy_mode"] == privacy
    assert payload["estimated_input_tokens"] >= 1
    if privacy == "share_safe":
        assert payload["privacy_lint"]["passed"] is True
        serialized = json.dumps(payload, ensure_ascii=False)
        assert str(FIXTURE_ROOT) not in serialized
        assert "alice@example.com" not in serialized
        assert "sk-fixture-placeholder" not in serialized
    else:
        assert payload["privacy_lint"] is None
    return payload


def _create_analysis(
    client: TestClient,
    packet_id: str,
    *,
    provider: str,
) -> dict[str, object]:
    response = client.post(
        "/v1/inventory/analyses",
        json={"packet_id": packet_id, "provider": provider},
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["analysis"]["packet_id"] == packet_id
    assert payload["analysis"]["provider"] == provider
    assert payload["receipt"]["network_calls"] == 0
    return payload


def _create_report(
    client: TestClient,
    *,
    snapshot_id: str,
    packet_id: str,
    analysis_id: str,
) -> dict[str, object]:
    response = client.post(
        "/v1/inventory/reports",
        json={
            "snapshot_id": snapshot_id,
            "packet_id": packet_id,
            "analysis_id": analysis_id,
        },
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["status"] == payload["manifest"]["analysis_status"]
    assert payload["manifest"]["snapshot_id"] == snapshot_id
    assert payload["manifest"]["packet_id"] == packet_id
    assert payload["manifest"]["analysis_id"] == analysis_id
    assert payload["download_url"].endswith(
        f"/v1/inventory/reports/{payload['report_id']}/download"
    )
    assert "path" not in json.dumps(payload).casefold()
    return payload


def test_inventory_api_rejects_real_scope_and_unsupported_options(settings: Settings) -> None:
    with TestClient(create_app(settings)) as client:
        real_scope = client.post(
            "/v1/inventory/scans",
            json={
                "fixture_id": "standard",
                "profile": "inventory_metadata",
                "paths": ["/Users/example/Documents"],
            },
        )
        assert real_scope.status_code == 422
        assert real_scope.json()["code"] == "REQUEST_VALIDATION_FAILED"

        wrong_fixture = client.post(
            "/v1/inventory/scans",
            json={"fixture_id": "other", "profile": "inventory_metadata"},
        )
        assert wrong_fixture.status_code == 422

        hashes = client.post(
            "/v1/inventory/scans",
            json={
                "fixture_id": "standard",
                "profile": "inventory_metadata",
                "compute_hash": True,
            },
        )
        assert hashes.status_code == 422
        assert hashes.json()["code"] == "REQUEST_VALIDATION_FAILED"


def test_inventory_api_none_and_fake_persist_without_id_collisions(
    settings: Settings,
    monkeypatch,
) -> None:
    before, before_hashes = _fixture_integrity()
    network = Mock(side_effect=AssertionError("network call is forbidden"))
    monkeypatch.setattr(socket, "create_connection", network)

    with TestClient(create_app(settings)) as client:
        scan = _create_scan(client)
        repeated_scan = _create_scan(client)
        assert repeated_scan["scan_id"] == scan["scan_id"]

        status_response = client.get(f"/v1/inventory/scans/{scan['scan_id']}")
        assert status_response.status_code == 200
        assert status_response.json() == scan

        snapshot = _create_snapshot(client, str(scan["scan_id"]))
        repeated_snapshot = _create_snapshot(client, str(scan["scan_id"]))
        assert repeated_snapshot["snapshot_id"] == snapshot["snapshot_id"]
        get_snapshot = client.get(
            f"/v1/inventory/snapshots/{snapshot['snapshot_id']}"
        )
        assert get_snapshot.status_code == 200
        assert get_snapshot.json() == snapshot

        local = _preview_packet(
            client,
            str(snapshot["snapshot_id"]),
            privacy="local_full",
            execution="none",
        )
        none = _create_analysis(
            client,
            str(local["packet"]["packet_id"]),
            provider="none",
        )
        assert none["analysis"]["status"] == "deterministic_only"
        assert none["receipt"]["provider_calls"] == 0
        none_report = _create_report(
            client,
            snapshot_id=str(snapshot["snapshot_id"]),
            packet_id=str(local["packet"]["packet_id"]),
            analysis_id=str(none["analysis"]["analysis_id"]),
        )

        safe = _preview_packet(
            client,
            str(snapshot["snapshot_id"]),
            privacy="share_safe",
            execution="local_only",
        )
        repeated_safe = _preview_packet(
            client,
            str(snapshot["snapshot_id"]),
            privacy="share_safe",
            execution="local_only",
        )
        assert repeated_safe["packet"] == safe["packet"]
        fake = _create_analysis(
            client,
            str(safe["packet"]["packet_id"]),
            provider="fake",
        )
        assert fake["analysis"]["status"] == "ai_complete"
        assert fake["receipt"]["provider_calls"] == 1
        fake_cached = _create_analysis(
            client,
            str(safe["packet"]["packet_id"]),
            provider="fake",
        )
        assert fake_cached["analysis"]["analysis_id"] == fake["analysis"]["analysis_id"]
        assert fake_cached["receipt"]["status"] == "cache_hit"
        assert fake_cached["receipt"]["provider_calls"] == 0
        fake_report = _create_report(
            client,
            snapshot_id=str(snapshot["snapshot_id"]),
            packet_id=str(safe["packet"]["packet_id"]),
            analysis_id=str(fake["analysis"]["analysis_id"]),
        )
        assert fake_report["report_id"] != none_report["report_id"]

        repeated_report = _create_report(
            client,
            snapshot_id=str(snapshot["snapshot_id"]),
            packet_id=str(safe["packet"]["packet_id"]),
            analysis_id=str(fake["analysis"]["analysis_id"]),
        )
        assert repeated_report == fake_report

        for report in (none_report, fake_report):
            report_id = report["report_id"]
            status_response = client.get(f"/v1/inventory/reports/{report_id}")
            assert status_response.status_code == 200
            assert status_response.json() == report
            download = client.get(f"/v1/inventory/reports/{report_id}/download")
            assert download.status_code == 200
            assert (
                download.headers["content-type"]
                == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            assert download.headers["content-disposition"].startswith("attachment;")
            assert download.headers["cache-control"] == "no-store"
            assert download.content.startswith(b"PK")
            assert (
                hashlib.sha256(download.content).hexdigest()
                == report["manifest"]["workbook_sha256"]
            )
            with ZipFile(BytesIO(download.content)) as archive:
                assert archive.testzip() is None
            workbook = load_workbook(BytesIO(download.content), read_only=True)
            assert workbook.sheetnames == SHEET_NAMES
            if report_id == fake_report["report_id"]:
                values = " ".join(
                    str(cell.value)
                    for sheet in workbook.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                    if cell.value is not None
                )
                assert "alice@example.com" not in values
                assert "sk-fixture-placeholder" not in values
                assert "IGNORE_PREVIOUS_INSTRUCTIONS" not in values
            workbook.close()

        analysis_get = client.get(
            f"/v1/inventory/analyses/{fake['analysis']['analysis_id']}"
        )
        assert analysis_get.status_code == 200
        assert analysis_get.json()["analysis"] == fake["analysis"]

        with sqlite3.connect(settings.database_path) as connection:
            packet_json = connection.execute(
                "SELECT payload_json FROM analysis_packets WHERE packet_id = ?",
                (safe["packet"]["packet_id"],),
            ).fetchone()[0]
        assert "alice@example.com" not in packet_json
        assert "sk-fixture-placeholder" not in packet_json
        assert str(FIXTURE_ROOT) not in packet_json

    after, after_hashes = _fixture_integrity()
    assert before == after
    assert before_hashes + after_hashes > 0
    assert network.call_count == 0


def test_inventory_api_persists_across_app_restart(settings: Settings) -> None:
    with TestClient(create_app(settings)) as client:
        scan = _create_scan(client)
        snapshot = _create_snapshot(client, str(scan["scan_id"]))
        packet = _preview_packet(
            client,
            str(snapshot["snapshot_id"]),
            privacy="share_safe",
            execution="local_only",
        )
        analysis = _create_analysis(
            client,
            str(packet["packet"]["packet_id"]),
            provider="fake",
        )
        report = _create_report(
            client,
            snapshot_id=str(snapshot["snapshot_id"]),
            packet_id=str(packet["packet"]["packet_id"]),
            analysis_id=str(analysis["analysis"]["analysis_id"]),
        )

    with TestClient(create_app(settings)) as restarted:
        assert restarted.get(
            f"/v1/inventory/scans/{scan['scan_id']}"
        ).status_code == 200
        assert restarted.get(
            f"/v1/inventory/snapshots/{snapshot['snapshot_id']}"
        ).status_code == 200
        assert restarted.get(
            f"/v1/inventory/analyses/{analysis['analysis']['analysis_id']}"
        ).status_code == 200
        cached = restarted.post(
            "/v1/inventory/analyses",
            json={
                "packet_id": packet["packet"]["packet_id"],
                "provider": "fake",
            },
        )
        assert cached.status_code == 201
        assert cached.json()["receipt"]["status"] == "cache_hit"
        assert cached.json()["receipt"]["provider_calls"] == 0
        assert restarted.get(
            f"/v1/inventory/reports/{report['report_id']}"
        ).status_code == 200
        assert restarted.get(
            f"/v1/inventory/reports/{report['report_id']}/download"
        ).status_code == 200


def test_failed_fake_analysis_is_retried_after_restart(settings: Settings) -> None:
    with TestClient(create_app(settings)) as client:
        scan = _create_scan(client)
        snapshot = _create_snapshot(client, str(scan["scan_id"]))
        packet = _preview_packet(
            client,
            str(snapshot["snapshot_id"]),
            privacy="share_safe",
            execution="local_only",
        )
        client.app.state.inventory_service.fake_provider.fail = True
        failed = _create_analysis(
            client,
            str(packet["packet"]["packet_id"]),
            provider="fake",
        )
        assert failed["analysis"]["status"] == "ai_failed_fallback"
        assert failed["receipt"]["status"] == "provider_failed"
        assert failed["receipt"]["provider_calls"] == 1

    with TestClient(create_app(settings)) as restarted:
        recovered = _create_analysis(
            restarted,
            str(packet["packet"]["packet_id"]),
            provider="fake",
        )
        assert recovered["analysis"]["status"] == "ai_complete"
        assert recovered["receipt"]["status"] == "succeeded"
        assert recovered["receipt"]["provider_calls"] == 1
        assert recovered["analysis"]["analysis_id"] != failed["analysis"]["analysis_id"]

        cached = _create_analysis(
            restarted,
            str(packet["packet"]["packet_id"]),
            provider="fake",
        )
        assert cached["analysis"]["analysis_id"] == recovered["analysis"]["analysis_id"]
        assert cached["receipt"]["status"] == "cache_hit"
        assert cached["receipt"]["provider_calls"] == 0


def test_inventory_api_returns_safe_not_found_errors(settings: Settings) -> None:
    with TestClient(create_app(settings)) as client:
        cases = [
            ("GET", "/v1/inventory/scans/scan_missing", None),
            ("GET", "/v1/inventory/snapshots/snapshot_missing", None),
            ("GET", "/v1/inventory/analyses/analysis_missing", None),
            ("GET", "/v1/inventory/reports/report_missing", None),
            ("GET", "/v1/inventory/reports/report_missing/download", None),
            ("POST", "/v1/inventory/snapshots", {"scan_id": "scan_missing"}),
            (
                "POST",
                "/v1/inventory/analysis-packets/preview",
                {
                    "snapshot_id": "snapshot_missing",
                    "artifact_privacy_mode": "share_safe",
                    "analysis_execution_mode": "local_only",
                },
            ),
            (
                "POST",
                "/v1/inventory/analyses",
                {"packet_id": "packet_missing", "provider": "none"},
            ),
            (
                "POST",
                "/v1/inventory/reports",
                {
                    "snapshot_id": "snapshot_missing",
                    "packet_id": "packet_missing",
                    "analysis_id": "analysis_missing",
                },
            ),
        ]
        for method, path, body in cases:
            response = client.request(method, path, json=body)
            assert response.status_code == 404
            serialized = json.dumps(response.json(), ensure_ascii=False)
            assert str(settings.database_path.parent) not in serialized


def test_inventory_api_blocks_mixed_lineage_and_unsupported_provider(
    settings: Settings,
) -> None:
    with TestClient(create_app(settings)) as client:
        scan = _create_scan(client)
        snapshot = _create_snapshot(client, str(scan["scan_id"]))
        local = _preview_packet(
            client,
            str(snapshot["snapshot_id"]),
            privacy="local_full",
            execution="none",
        )
        none = _create_analysis(
            client,
            str(local["packet"]["packet_id"]),
            provider="none",
        )
        safe = _preview_packet(
            client,
            str(snapshot["snapshot_id"]),
            privacy="share_safe",
            execution="local_only",
        )
        fake = _create_analysis(
            client,
            str(safe["packet"]["packet_id"]),
            provider="fake",
        )

        mismatch = client.post(
            "/v1/inventory/reports",
            json={
                "snapshot_id": snapshot["snapshot_id"],
                "packet_id": local["packet"]["packet_id"],
                "analysis_id": fake["analysis"]["analysis_id"],
            },
        )
        assert mismatch.status_code == 409
        assert mismatch.json()["code"] == "REPORT_LINEAGE_MISMATCH"

        unsupported = client.post(
            "/v1/inventory/analyses",
            json={
                "packet_id": safe["packet"]["packet_id"],
                "provider": "openai",
            },
        )
        assert unsupported.status_code == 422
        assert unsupported.json()["code"] == "REQUEST_VALIDATION_FAILED"

        wrong_mode = client.post(
            "/v1/inventory/analyses",
            json={
                "packet_id": local["packet"]["packet_id"],
                "provider": "fake",
            },
        )
        assert wrong_mode.status_code == 409
        assert wrong_mode.json()["code"] == "ANALYSIS_MODE_MISMATCH"

        assert none["receipt"]["provider_calls"] == 0


def test_inventory_download_blocks_missing_and_hash_mismatched_artifacts(
    settings: Settings,
    monkeypatch,
    tmp_path: Path,
) -> None:
    with TestClient(create_app(settings)) as client:
        scan = _create_scan(client)
        snapshot = _create_snapshot(client, str(scan["scan_id"]))
        safe = _preview_packet(
            client,
            str(snapshot["snapshot_id"]),
            privacy="share_safe",
            execution="local_only",
        )
        fake = _create_analysis(
            client,
            str(safe["packet"]["packet_id"]),
            provider="fake",
        )
        report = _create_report(
            client,
            snapshot_id=str(snapshot["snapshot_id"]),
            packet_id=str(safe["packet"]["packet_id"]),
            analysis_id=str(fake["analysis"]["analysis_id"]),
        )
        with monkeypatch.context() as hash_guard:
            hash_guard.setattr(
                inventory_service_module,
                "workbook_sha256",
                lambda _path: "0" * 64,
            )
            mismatch = client.get(
                f"/v1/inventory/reports/{report['report_id']}/download"
            )
        assert mismatch.status_code == 409
        assert mismatch.json()["code"] == "REPORT_ARTIFACT_HASH_MISMATCH"

    alternate = Settings(
        database_path=settings.database_path,
        inventory_report_path=tmp_path / "missing-artifacts",
        allowed_origins=settings.allowed_origins,
    )
    with TestClient(create_app(alternate)) as client:
        missing = client.get(
            f"/v1/inventory/reports/{report['report_id']}/download"
        )
    assert missing.status_code == 409
    assert missing.json()["code"] == "REPORT_ARTIFACT_MISSING"
    assert str(settings.resolved_inventory_report_path) not in json.dumps(
        missing.json(),
        ensure_ascii=False,
    )


def test_inventory_product_boundary_has_zero_source_reads_hashes_network_or_actions(
    settings: Settings,
    monkeypatch,
) -> None:
    before, before_hashes = _fixture_integrity()
    standard_manifest = FIXTURE_ROOT / "standard" / "manifest.json"
    original_open = Path.open
    original_unlink = Path.unlink
    original_rename = Path.rename
    original_replace = Path.replace
    content_reads = Mock(side_effect=AssertionError("fixture content read is forbidden"))
    file_actions = Mock(side_effect=AssertionError("fixture file action is forbidden"))
    network = Mock(side_effect=AssertionError("network call is forbidden"))
    product_hash = Mock(side_effect=AssertionError("product file hash is forbidden"))

    def guarded_open(path: Path, *args, **kwargs):
        mode = args[0] if args else kwargs.get("mode", "r")
        if (
            path != standard_manifest
            and path.is_relative_to(FIXTURE_ROOT)
            and "r" in mode
        ):
            content_reads(path)
        return original_open(path, *args, **kwargs)

    def guarded_unlink(path: Path, *args, **kwargs):
        if path.is_relative_to(FIXTURE_ROOT):
            file_actions(path)
        return original_unlink(path, *args, **kwargs)

    def guarded_rename(path: Path, target, *args, **kwargs):
        if path.is_relative_to(FIXTURE_ROOT):
            file_actions(path)
        return original_rename(path, target, *args, **kwargs)

    def guarded_replace(path: Path, target, *args, **kwargs):
        if path.is_relative_to(FIXTURE_ROOT):
            file_actions(path)
        return original_replace(path, target, *args, **kwargs)

    with monkeypatch.context() as boundary:
        boundary.setattr(Path, "open", guarded_open)
        boundary.setattr(Path, "unlink", guarded_unlink)
        boundary.setattr(Path, "rename", guarded_rename)
        boundary.setattr(Path, "replace", guarded_replace)
        boundary.setattr(socket, "create_connection", network)
        boundary.setattr(p0_hashing, "partial_sha256", product_hash)
        boundary.setattr(p0_hashing, "full_sha256", product_hash)

        with TestClient(create_app(settings)) as client:
            scan = _create_scan(client)
            snapshot = _create_snapshot(client, str(scan["scan_id"]))
            safe = _preview_packet(
                client,
                str(snapshot["snapshot_id"]),
                privacy="share_safe",
                execution="local_only",
            )
            fake = _create_analysis(
                client,
                str(safe["packet"]["packet_id"]),
                provider="fake",
            )
            report = _create_report(
                client,
                snapshot_id=str(snapshot["snapshot_id"]),
                packet_id=str(safe["packet"]["packet_id"]),
                analysis_id=str(fake["analysis"]["analysis_id"]),
            )
            assert client.get(
                f"/v1/inventory/reports/{report['report_id']}/download"
            ).status_code == 200

    after, after_hashes = _fixture_integrity()
    assert before == after
    assert before_hashes + after_hashes > 0
    assert content_reads.call_count == 0
    assert file_actions.call_count == 0
    assert product_hash.call_count == 0
    assert network.call_count == 0
