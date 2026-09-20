from __future__ import annotations

from openpyxl import load_workbook
import pytest

from core.analysis.orchestrator import analyze_packet
from core.analysis.packet import build_packet
from core.analysis.providers.fake import FakeProvider
from core.analysis.schemas import AnalysisConsentV1, default_analysis_budget
from core.inventory.models import canonical_json
from core.reports.models import SHEET_NAMES
from core.reports.validation import (
    WorkbookValidationError,
    build_report_manifest,
    validate_workbook,
)
from core.reports.view_model import build_report_view_model
from core.reports.workbook import render_workbook
from core.tests._inventory_helpers import make_snapshot


def _render(tmp_path, provider: str, privacy: str):
    snapshot, _ = make_snapshot(tmp_path)
    packet = build_packet(
        snapshot,
        artifact_privacy_mode=privacy,
        analysis_execution_mode="none" if provider == "none" else "local_only",
    )
    outcome = analyze_packet(
        packet,
        provider=provider,
        gateway=FakeProvider() if provider == "fake" else None,
    )
    view = build_report_view_model(
        snapshot,
        packet,
        outcome.analysis,
        outcome.receipt,
    )
    path = tmp_path / f"{provider}-{privacy}.xlsx"
    render_workbook(view, path)
    validation = validate_workbook(
        path,
        view=view,
        snapshot=snapshot,
        packet=packet,
        analysis=outcome.analysis,
    )
    manifest = build_report_manifest(path=path, view=view, validation=validation)
    return path, manifest


def test_none_and_fake_workbooks_reopen_with_fixed_five_sheets(tmp_path) -> None:
    none_path, none_manifest = _render(tmp_path / "none", "none", "local_full")
    fake_path, fake_manifest = _render(tmp_path / "fake", "fake", "share_safe")

    for path, manifest in ((none_path, none_manifest), (fake_path, fake_manifest)):
        workbook = load_workbook(path, data_only=False)
        assert workbook.sheetnames == SHEET_NAMES
        assert all(sheet.freeze_panes == "A2" for sheet in workbook.worksheets)
        assert all(sheet.auto_filter.ref for sheet in workbook.worksheets)
        assert not any(sheet.merged_cells.ranges for sheet in workbook.worksheets)
        assert workbook["空间体检概览"]["B4"].is_date
        assert manifest.workbook_sha256
        workbook.close()


def test_share_safe_workbook_does_not_contain_fixture_path(tmp_path) -> None:
    path, _ = _render(tmp_path / "fake", "fake", "share_safe")
    workbook = load_workbook(path, data_only=False)
    values = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    workbook.close()

    assert str(tmp_path) not in values
    assert "alice@example.com" not in values
    assert "已核验重复" not in values
    assert "尚未核验" in values


def test_synthetic_workbook_audit_contains_consent_and_zero_network(
    tmp_path,
) -> None:
    snapshot, _ = make_snapshot(tmp_path)
    packet = build_packet(
        snapshot,
        artifact_privacy_mode="share_safe",
        analysis_execution_mode="cloud_metadata_minimized",
    )
    packet_bytes = len(canonical_json(packet).encode("utf-8"))
    budget = default_analysis_budget()
    consent = AnalysisConsentV1(
        consent_receipt_id=(
            "consent_receipt_00000000-0000-4000-8000-000000000100"
        ),
        client_action_id=(
            "client_action_00000000-0000-4000-8000-000000000100"
        ),
        confirmed=True,
        snapshot_id=snapshot.snapshot_id,
        snapshot_sha256=snapshot.snapshot_sha256,
        packet_id=packet.packet_id,
        packet_sha256=packet.packet_sha256,
        packet_bytes=packet_bytes,
        artifact_privacy_mode="share_safe",
        analysis_execution_mode="cloud_metadata_minimized",
        provider="fake",
        model="fixture-analyst-v1",
        prompt_version="report-analysis-v1",
        consent_schema_version="1.0",
        analysis_schema_version="ai-analysis-v1",
        disclosure_version="synthetic-disclosure-v1",
        budget=budget.model_dump(),
        cost_basis="synthetic_zero_external_cost",
    )
    outcome = analyze_packet(
        packet,
        provider="fake",
        gateway=FakeProvider(),
        budget=budget,
        consent=consent,
    )
    view = build_report_view_model(
        snapshot,
        packet,
        outcome.analysis,
        outcome.receipt,
    )
    path = render_workbook(view, tmp_path / "synthetic.xlsx")
    workbook = load_workbook(path, data_only=False)
    audit = {
        row[0].value: row[1].value
        for row in workbook["报告说明与记录"].iter_rows(min_row=2, max_col=2)
    }
    workbook.close()

    assert audit["analysis_execution_mode"] == "cloud_metadata_minimized"
    assert audit["consent_receipt_id"] == (
        "consent_receipt_00000000-0000-4000-8000-000000000100"
    )
    assert audit["provider_calls"] == 1
    assert audit["network_calls"] == 0
    assert audit["cost_basis"] == "synthetic_zero_external_cost"

    unsafe_view = view.model_copy(
        update={
            "consent_receipt_id": "=1+1",
            "analysis_error_code": "@SUM(A1:A2)",
        }
    )
    unsafe_path = render_workbook(unsafe_view, tmp_path / "formula-like.xlsx")
    unsafe_workbook = load_workbook(unsafe_path, data_only=False)
    unsafe_audit = {
        row[0].value: row[1]
        for row in unsafe_workbook["报告说明与记录"].iter_rows(
            min_row=2,
            max_col=2,
        )
    }
    assert unsafe_audit["consent_receipt_id"].data_type != "f"
    assert unsafe_audit["consent_receipt_id"].value == "'=1+1"
    assert unsafe_audit["analysis_error_code"].data_type != "f"
    assert unsafe_audit["analysis_error_code"].value == "'@SUM(A1:A2)"
    assert not any(
        cell.data_type == "f"
        for sheet in unsafe_workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )
    unsafe_workbook.close()

    injected_workbook = load_workbook(unsafe_path, data_only=False)
    injected_workbook["报告说明与记录"]["B2"] = "=1+1"
    injected_workbook.save(unsafe_path)
    injected_workbook.close()
    with pytest.raises(WorkbookValidationError, match="FORMULA_CELL"):
        validate_workbook(
            unsafe_path,
            view=unsafe_view,
            snapshot=snapshot,
            packet=packet,
            analysis=outcome.analysis,
        )
