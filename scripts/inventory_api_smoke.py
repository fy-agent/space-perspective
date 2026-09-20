from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import socket
import sys
from unittest.mock import patch
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from core.app.config import Settings
from core.app.main import create_app
from core.reports.models import SHEET_NAMES

FIXTURE_ROOT = ROOT / "fixtures" / "ai-report"
TRIALS_ROOT = ROOT / "output" / "inventory-api-trials"


def _fixture_integrity() -> tuple[list[dict[str, object]], int]:
    records: list[dict[str, object]] = []
    hashes = 0
    for path in sorted(
        FIXTURE_ROOT.rglob("*"),
        key=lambda item: item.relative_to(FIXTURE_ROOT).as_posix(),
    ):
        info = path.stat(follow_symlinks=False)
        record: dict[str, object] = {
            "relative_path": path.relative_to(FIXTURE_ROOT).as_posix(),
            "kind": "directory" if path.is_dir() else "file",
            "size": info.st_size,
            "mtime_ns": info.st_mtime_ns,
            "sha256": None,
        }
        if path.is_file():
            digest = hashlib.sha256()
            with path.open("rb") as source:
                while chunk := source.read(1024 * 1024):
                    digest.update(chunk)
            record["sha256"] = digest.hexdigest()
            hashes += 1
        records.append(record)
    return records, hashes


def _post(client: TestClient, path: str, payload: dict[str, object]) -> dict[str, object]:
    response = client.post(path, json=payload)
    if response.status_code not in {200, 201}:
        raise RuntimeError(f"{path} failed: {response.status_code} {response.text}")
    return response.json()


def _packet(
    client: TestClient,
    snapshot_id: str,
    *,
    privacy: str,
    execution: str,
) -> dict[str, object]:
    return _post(
        client,
        "/v1/inventory/analysis-packets/preview",
        {
            "snapshot_id": snapshot_id,
            "artifact_privacy_mode": privacy,
            "analysis_execution_mode": execution,
            "top_k": 200,
        },
    )


def _analysis(
    client: TestClient,
    packet_id: str,
    provider: str,
) -> dict[str, object]:
    return _post(
        client,
        "/v1/inventory/analyses",
        {"packet_id": packet_id, "provider": provider},
    )


def _report(
    client: TestClient,
    *,
    snapshot_id: str,
    packet_id: str,
    analysis_id: str,
) -> dict[str, object]:
    return _post(
        client,
        "/v1/inventory/reports",
        {
            "snapshot_id": snapshot_id,
            "packet_id": packet_id,
            "analysis_id": analysis_id,
        },
    )


def main() -> int:
    trial_id = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H-%M-%SZ-") + uuid4().hex[:12]
    output = TRIALS_ROOT / trial_id
    output.mkdir(parents=True, exist_ok=False)
    before, before_hashes = _fixture_integrity()
    settings = Settings(
        database_path=output / "inventory-api.sqlite3",
        inventory_report_path=output / "reports",
    )
    network_guard = patch.object(
        socket,
        "create_connection",
        side_effect=AssertionError("network call is forbidden"),
    )
    with network_guard, TestClient(create_app(settings)) as client:
        scan = _post(
            client,
            "/v1/inventory/scans",
            {
                "fixture_id": "standard",
                "profile": "inventory_metadata",
                "compute_hash": False,
            },
        )
        snapshot = _post(
            client,
            "/v1/inventory/snapshots",
            {"scan_id": scan["scan_id"]},
        )
        local = _packet(
            client,
            snapshot["snapshot_id"],
            privacy="local_full",
            execution="none",
        )
        none = _analysis(client, local["packet"]["packet_id"], "none")
        none_report = _report(
            client,
            snapshot_id=snapshot["snapshot_id"],
            packet_id=local["packet"]["packet_id"],
            analysis_id=none["analysis"]["analysis_id"],
        )
        safe = _packet(
            client,
            snapshot["snapshot_id"],
            privacy="share_safe",
            execution="local_only",
        )
        fake = _analysis(client, safe["packet"]["packet_id"], "fake")
        fake_cached = _analysis(client, safe["packet"]["packet_id"], "fake")
        fake_report = _report(
            client,
            snapshot_id=snapshot["snapshot_id"],
            packet_id=safe["packet"]["packet_id"],
            analysis_id=fake["analysis"]["analysis_id"],
        )

        reports: list[dict[str, object]] = []
        for report in (none_report, fake_report):
            response = client.get(
                f"/v1/inventory/reports/{report['report_id']}/download"
            )
            if response.status_code != 200:
                raise RuntimeError(f"download failed: {response.status_code} {response.text}")
            digest = hashlib.sha256(response.content).hexdigest()
            if digest != report["manifest"]["workbook_sha256"]:
                raise RuntimeError("download hash does not match manifest")
            workbook = load_workbook(
                settings.inventory_report_path
                / report["report_id"]
                / report["manifest"]["filename"],
                read_only=True,
            )
            if workbook.sheetnames != SHEET_NAMES:
                raise RuntimeError("workbook sheets do not match frozen contract")
            workbook.close()
            reports.append(
                {
                    "report_id": report["report_id"],
                    "provider": report["manifest"]["provider"],
                    "analysis_status": report["manifest"]["analysis_status"],
                    "workbook_sha256": digest,
                    "path": str(
                        settings.inventory_report_path
                        / report["report_id"]
                        / report["manifest"]["filename"]
                    ),
                }
            )

    after, after_hashes = _fixture_integrity()
    receipt = {
        "status": "passed" if before == after else "failed",
        "evidence_level": "fixture_e2e",
        "trial_directory": str(output),
        "scan_id": scan["scan_id"],
        "snapshot_id": snapshot["snapshot_id"],
        "reports": reports,
        "none_provider_calls": none["receipt"]["provider_calls"],
        "fake_provider_calls": fake["receipt"]["provider_calls"],
        "fake_cache_hit_provider_calls": fake_cached["receipt"]["provider_calls"],
        "fake_cache_status": fake_cached["receipt"]["status"],
        "product_content_reads": scan["safety"]["product_content_reads"],
        "product_file_hashes": scan["safety"]["product_file_hashes"],
        "network_calls": scan["safety"]["network_calls"],
        "file_actions": scan["safety"]["file_actions"],
        "fixture_integrity_hashes": before_hashes + after_hashes,
        "fixture_unchanged": before == after,
        "share_safe_privacy_lint": safe["privacy_lint"],
    }
    receipt_path = output / "api-e2e-receipt.json"
    receipt_path.write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(receipt, ensure_ascii=False))
    return 0 if receipt["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
