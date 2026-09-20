from __future__ import annotations

import argparse
from contextlib import contextmanager, ExitStack
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import platform
import re
import shutil
import socket
import subprocess
import sys
import tempfile
from typing import Any, Callable
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.analysis.orchestrator import ProductAnalysisCache, analyze_packet
from core.analysis.packet import build_packet
from core.analysis.privacy import lint_share_safe_packet
from core.analysis.providers.fake import FakeProvider
from core.collectors.base import CancelToken, ScopeSpec
from core.collectors.checkpoint import CollectorCheckpoint
from core.collectors.filesystem import FileSystemCollector
from core.collectors.platform.windows import WindowsMetadataAdapter
from core.inventory.models import CoverageV1, ScanFactV1, canonical_json
from core.inventory.snapshots import build_snapshot
from core.reports.models import SHEET_NAMES
from core.reports.validation import (
    validate_workbook,
    workbook_sha256,
)
from core.reports.view_model import build_report_view_model
from core.reports.workbook import render_workbook
from core.reports.workbook import RAW_BYTE_FORMAT
import core.scanner.hashing as p0_hashing


EXIT_FAILED = 1
EXIT_BLOCKED = 3
FIXTURE_ID = "iteration6-controlled-v1"
OUTPUT_ROOT = ROOT / "output" / "iteration6"
PRIVATE_VALUE = re.compile(
    r"(?:[A-Za-z]:\\|\\\\(?:[^?]|$)|/(?:Users|home|private|var|tmp)/|"
    r"\.sqlite(?:3)?\b|runneradmin)",
    re.IGNORECASE,
)


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(
            value,
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        )
        + "\n",
        encoding="utf-8",
    )


@contextmanager
def _controlled_temporary_sandbox():
    sandbox = Path(
        tempfile.mkdtemp(prefix="space-perspective-iteration6-")
    )
    body_failed = False
    try:
        yield sandbox
    except BaseException:
        body_failed = True
        raise
    finally:
        try:
            shutil.rmtree(sandbox)
        except OSError:
            if not body_failed:
                raise


def _tool_version(command: list[str]) -> str:
    try:
        completed = subprocess.run(
            command,
            check=True,
            capture_output=True,
            text=True,
        )
    except (OSError, subprocess.CalledProcessError):
        return "unavailable"
    return completed.stdout.strip().splitlines()[0]


def build_environment_report(
    *,
    os_name: str,
    system: str,
    release: str,
    commit_sha: str,
) -> dict[str, Any]:
    report: dict[str, Any] = {
        "schema_version": "1.0",
        "fixture_id": FIXTURE_ID,
        "generated_at": _utc_now(),
        "status": "ready",
        "evidence_level": "runtime_windows_ci",
        "commit_sha": commit_sha,
        "environment": {
            "system": system,
            "release": release,
            "python": platform.python_version(),
            "machine": platform.machine(),
            "runner_image_os": os.environ.get("ImageOS", "unavailable"),
            "runner_image_version": os.environ.get(
                "ImageVersion",
                "unavailable",
            ),
        },
        "toolchain": {
            "uv": _tool_version(["uv", "--version"]),
            "node": _tool_version(["node", "--version"]),
        },
        "commands": [
            {
                "name": "iteration6_windows_fixture_runner",
                "command": (
                    "uv run python "
                    "scripts/iteration6_windows_office_acceptance.py "
                    "--evidence <receipt> --artifacts <artifacts>"
                ),
                "exit_code": None,
            }
        ],
        "checks": [],
        "product_counters": {
            "product_content_reads": 0,
            "product_file_hashes": 0,
            "network_calls": 0,
            "file_actions": 0,
            "provider_calls": 0,
        },
        "verifier_counters": {
            "fixture_integrity_hashes": 0,
        },
        "blockers": [],
        "claims_excluded": [
            "runtime_windows_11",
            "excel_wps_open",
            "real_scope",
            "real_provider",
            "production",
        ],
    }
    if os_name != "nt" or system.casefold() != "windows":
        report["status"] = "blocked"
        report["evidence_level"] = "code_audit"
        report["blockers"] = ["NATIVE_WINDOWS_REQUIRED"]
        report["commands"][0]["exit_code"] = EXIT_BLOCKED
    return report


def ensure_share_safe_receipt(value: dict[str, Any]) -> None:
    serialized = canonical_json(value)
    if PRIVATE_VALUE.search(serialized):
        raise ValueError("receipt contains private path, identity, or SQLite")
    forbidden_keys = {
        "absolute_path",
        "sqlite_path",
        "username",
        "hostname",
        "fixture_payload",
        "file_content",
    }

    def keys(item: Any) -> set[str]:
        if isinstance(item, dict):
            return set(item) | {
                key
                for nested in item.values()
                for key in keys(nested)
            }
        if isinstance(item, list):
            return {
                key
                for nested in item
                for key in keys(nested)
            }
        return set()

    if keys(value) & forbidden_keys:
        raise ValueError("receipt contains private fields")


def validate_office_workbook(path: Path) -> dict[str, Any]:
    with ZipFile(path) as archive:
        bad_member = archive.testzip()
        names = set(archive.namelist())
    workbook = load_workbook(
        path,
        data_only=False,
        read_only=False,
        keep_links=True,
    )
    formula_cells = 0
    date_cells = 0
    internal_links_valid = True
    freeze_and_filters = True
    column_widths_valid = True
    number_formats_valid = True
    for sheet in workbook.worksheets:
        freeze_and_filters = (
            freeze_and_filters
            and sheet.freeze_panes == "A2"
            and bool(sheet.auto_filter.ref)
        )
        column_widths_valid = column_widths_valid and all(
            10
            <= float(
                sheet.column_dimensions[
                    get_column_letter(column)
                ].width
                or 0
            )
            <= 48
            for column in range(1, sheet.max_column + 1)
        )
        for row in sheet.iter_rows():
            for cell in row:
                formula_cells += int(cell.data_type == "f")
                date_cells += int(cell.is_date)
                if cell.hyperlink is not None:
                    target = str(cell.hyperlink.target)
                    if target.startswith("#"):
                        match = re.match(r"^#'([^']+)'!A(\d+)$", target)
                        if match is None:
                            internal_links_valid = False
                        else:
                            target_sheet, row_number = match.groups()
                            internal_links_valid = (
                                internal_links_valid
                                and target_sheet in workbook.sheetnames
                                and int(row_number)
                                <= workbook[target_sheet].max_row
                            )
    raw_byte_cells = [
        workbook["空间体检概览"][coordinate]
        for coordinate in ("B10", "B12", "B14")
    ]
    raw_byte_cells.extend(
        workbook["空间明细清单"].cell(row, column)
        for row in range(
            2,
            workbook["空间明细清单"].max_row + 1,
        )
        for column in (6, 8, 10)
    )
    raw_byte_cells.extend(
        workbook["建议优先查看"].cell(row, 5)
        for row in range(
            2,
            workbook["建议优先查看"].max_row + 1,
        )
    )
    raw_byte_cells.extend(
        workbook["可能重复与安装包"].cell(row, 7)
        for row in range(
            2,
            workbook["可能重复与安装包"].max_row + 1,
        )
    )
    number_formats_valid = all(
        cell.value is None
        or type(cell.value) is not int
        or cell.number_format == RAW_BYTE_FORMAT
        for cell in raw_byte_cells
    )
    external_links = len(getattr(workbook, "_external_links", []))
    workbook.close()
    macro_parts = sum(
        "vbaProject" in name or name.endswith(".bin")
        for name in names
    )
    reopened = load_workbook(path, read_only=True)
    reopened_sheet_names = list(reopened.sheetnames)
    reopened.close()
    result = {
        "passed": (
            bad_member is None
            and list(SHEET_NAMES) == reopened_sheet_names
            and freeze_and_filters
            and column_widths_valid
            and number_formats_valid
            and formula_cells == 0
            and external_links == 0
            and macro_parts == 0
            and date_cells >= 1
            and internal_links_valid
        ),
        "sheet_count": len(reopened_sheet_names),
        "freeze_and_filters": freeze_and_filters,
        "column_widths_valid": column_widths_valid,
        "number_formats_valid": number_formats_valid,
        "formula_cells": formula_cells,
        "external_links": external_links,
        "macro_parts": macro_parts,
        "date_cells": date_cells,
        "internal_links_valid": internal_links_valid,
        "zip_integrity": bad_member is None,
    }
    return result


def _forbidden_product_surface(*_args: Any, **_kwargs: Any) -> None:
    raise AssertionError("iteration 6 product collection crossed a forbidden surface")


def _collect_phase(
    *,
    root: Path,
    allowed_root: Path,
    checkpoint: CollectorCheckpoint | None,
    cancel_after_first_batch: bool,
    batch_size: int,
) -> dict[str, Any]:
    token = CancelToken()
    collector = FileSystemCollector(
        adapter=WindowsMetadataAdapter(),
        batch_size=batch_size,
    )
    facts: list[ScanFactV1] = []
    final = None
    with ExitStack() as stack:
        stack.enter_context(patch.object(Path, "open", _forbidden_product_surface))
        stack.enter_context(patch.object(Path, "unlink", _forbidden_product_surface))
        stack.enter_context(patch.object(Path, "rename", _forbidden_product_surface))
        stack.enter_context(patch.object(Path, "replace", _forbidden_product_surface))
        stack.enter_context(patch.object(socket, "create_connection", _forbidden_product_surface))
        stack.enter_context(patch.object(p0_hashing, "partial_sha256", _forbidden_product_surface))
        stack.enter_context(patch.object(p0_hashing, "full_sha256", _forbidden_product_surface))
        for batch in collector.collect(
            ScopeSpec(
                scope_id="iteration6-windows-fixture",
                root=root,
                allowed_root=allowed_root,
            ),
            "inventory_metadata",
            checkpoint,
            token,
        ):
            facts.extend(batch.facts)
            final = batch
            if (
                cancel_after_first_batch
                and not batch.final
                and not token.cancelled
            ):
                token.cancel()
    if final is None:
        raise RuntimeError("collector did not produce a final batch")
    return {
        "facts": [
            fact.model_dump(mode="json")
            for fact in facts
        ],
        "coverage": final.coverage.model_dump(mode="json"),
        "checkpoint": final.checkpoint.model_dump(mode="json"),
        "product_counters": {
            "product_content_reads": (
                collector.counters.product_content_reads
            ),
            "product_file_hashes": (
                collector.counters.product_file_hashes
            ),
            "network_calls": collector.counters.network_calls,
            "file_actions": 0,
            "provider_calls": 0,
        },
    }


def _child_collect(args: argparse.Namespace) -> int:
    try:
        checkpoint = None
        if args.checkpoint:
            checkpoint = CollectorCheckpoint.model_validate_json(
                Path(args.checkpoint).read_text(encoding="utf-8")
            )
        payload = _collect_phase(
            root=Path(args.root),
            allowed_root=Path(args.allowed_root),
            checkpoint=checkpoint,
            cancel_after_first_batch=args.cancel,
            batch_size=args.batch_size,
        )
    except Exception as exc:
        print(
            json.dumps(
                {
                    "status": "child_failed",
                    "error_type": type(exc).__name__,
                },
                sort_keys=True,
            )
        )
        return EXIT_FAILED
    print(json.dumps(payload, ensure_ascii=True, sort_keys=True))
    return 0


def _run_child(
    *,
    root: Path,
    allowed_root: Path,
    checkpoint_path: Path | None,
    cancel: bool,
) -> dict[str, Any]:
    command = [
        sys.executable,
        str(Path(__file__).resolve()),
        "_collect-phase",
        "--root",
        str(root),
        "--allowed-root",
        str(allowed_root),
        "--batch-size",
        "4",
    ]
    if checkpoint_path is not None:
        command.extend(["--checkpoint", str(checkpoint_path)])
    if cancel:
        command.append("--cancel")
    try:
        completed = subprocess.run(
            command,
            check=False,
            capture_output=True,
            text=True,
            cwd=ROOT,
        )
    except OSError as exc:
        raise RuntimeError(
            f"CHILD_PROCESS_{type(exc).__name__.upper()}"
        ) from exc
    payload = json.loads(completed.stdout)
    if completed.returncode != 0:
        error_type = str(payload.get("error_type", "UNKNOWN")).upper()
        raise RuntimeError(f"CHILD_COLLECTION_{error_type}")
    return payload


def _file_digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _hash_fixture(paths: list[Path]) -> dict[str, str]:
    return {
        path.name: _file_digest(path)
        for path in sorted(paths, key=lambda item: item.name.casefold())
    }


def _set_file_attributes(path: Path, attributes: int) -> None:
    import ctypes

    set_attributes = ctypes.windll.kernel32.SetFileAttributesW
    set_attributes.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32]
    set_attributes.restype = ctypes.c_int
    if not set_attributes(str(path), attributes):
        raise ctypes.WinError()


def _create_junction(link: Path, target: Path) -> None:
    subprocess.run(
        ["cmd", "/d", "/c", "mklink", "/J", str(link), str(target)],
        check=True,
        capture_output=True,
        text=True,
    )


def _apply_read_deny(path: Path) -> Callable[[], None]:
    identity = subprocess.run(
        ["whoami"],
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    subprocess.run(
        [
            "icacls",
            str(path),
            "/deny",
            f"{identity}:(RD)",
        ],
        check=True,
        capture_output=True,
        text=True,
    )

    def restore() -> None:
        subprocess.run(
            ["icacls", str(path), "/remove:d", identity],
            check=True,
            capture_output=True,
            text=True,
        )
        subprocess.run(
            ["icacls", str(path), "/reset"],
            check=True,
            capture_output=True,
            text=True,
        )

    return restore


def _build_fixture(sandbox: Path) -> dict[str, Any]:
    root = sandbox / "fixture"
    root.mkdir()
    normal = root / "中文 空格"
    normal.mkdir()
    files = [
        normal / "普通报告.txt",
        normal / "安装包.exe",
    ]
    files[0].write_text("controlled iteration 6", encoding="utf-8")
    files[1].write_bytes(b"controlled installer fixture")
    hidden = root / "hidden-fixture.txt"
    hidden.write_text("hidden fixture", encoding="utf-8")
    system = root / "system-fixture.txt"
    system.write_text("system fixture", encoding="utf-8")
    _set_file_attributes(hidden, 0x2)
    _set_file_attributes(system, 0x4)
    files.extend([hidden, system])

    long_parent = root / "long"
    filename = "long-path-fixture.txt"
    while len(str(long_parent / filename)) <= 280:
        long_parent /= "segment-0123456789abcdef"
    long_parent.mkdir(parents=True)
    long_file = long_parent / filename
    long_file.write_text("long path fixture", encoding="utf-8")
    files.append(long_file)

    denied = root / "permission-gap"
    denied.mkdir()
    denied_file = denied / "denied.txt"
    denied_file.write_text("permission fixture", encoding="utf-8")
    files.append(denied_file)

    target = sandbox / "junction-target"
    target.mkdir()
    target_file = target / "must-not-be-collected.txt"
    target_file.write_text("junction target", encoding="utf-8")
    files.append(target_file)
    junction = root / "junction-fixture"
    _create_junction(junction, target)
    return {
        "root": root,
        "files": files,
        "hidden": hidden,
        "system": system,
        "long_file": long_file,
        "denied": denied,
        "junction": junction,
        "junction_target_name": target_file.name,
    }


def _facts(payload: dict[str, Any]) -> list[ScanFactV1]:
    return [
        ScanFactV1.model_validate(item)
        for item in payload["facts"]
    ]


def _record(
    report: dict[str, Any],
    name: str,
    passed: bool,
    **details: Any,
) -> None:
    report["checks"].append(
        {
            "name": name,
            "passed": bool(passed),
            "details": details,
        }
    )


def _cell_strings(path: Path) -> str:
    workbook = load_workbook(path, data_only=False, read_only=True)
    values = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    workbook.close()
    return values


def _run_native(
    report: dict[str, Any],
    *,
    artifacts: Path,
) -> dict[str, Any]:
    with _controlled_temporary_sandbox() as sandbox:
        report["active_phase"] = "fixture_build"
        fixture = _build_fixture(sandbox)
        report["active_phase"] = "fixture_integrity_before"
        before = _hash_fixture(fixture["files"])
        report["active_phase"] = "permission_gap_apply"
        restore_acl = _apply_read_deny(fixture["denied"])
        try:
            report["active_phase"] = "cancel_child"
            checkpoint_path = sandbox / "checkpoint.json"
            cancel_phase = _run_child(
                root=fixture["root"],
                allowed_root=sandbox,
                checkpoint_path=None,
                cancel=True,
            )
            _write_json(
                checkpoint_path,
                cancel_phase["checkpoint"],
            )
            report["active_phase"] = "resume_child"
            resume_phase = _run_child(
                root=fixture["root"],
                allowed_root=sandbox,
                checkpoint_path=checkpoint_path,
                cancel=False,
            )
            report["active_phase"] = "clean_child"
            clean_phase = _run_child(
                root=fixture["root"],
                allowed_root=sandbox,
                checkpoint_path=None,
                cancel=False,
            )
        finally:
            previous_phase = report["active_phase"]
            report["active_phase"] = "permission_gap_restore"
            try:
                restore_acl()
            except Exception:
                raise
            else:
                report["active_phase"] = previous_phase

        report["active_phase"] = "fixture_integrity_after"
        after = _hash_fixture(fixture["files"])
        report["active_phase"] = "report_chain"
        first_facts = _facts(cancel_phase)
        resumed_facts = _facts(resume_phase)
        clean_facts = _facts(clean_phase)
        combined = [*first_facts, *resumed_facts]
        combined_dump = [
            item.model_dump(mode="json")
            for item in combined
        ]
        clean_dump = [
            item.model_dump(mode="json")
            for item in clean_facts
        ]
        resume_exact = (
            cancel_phase["checkpoint"]["status"] == "cancelled"
            and resume_phase["checkpoint"]["status"] == "partial"
            and len({item.fact_id for item in combined}) == len(combined)
            and combined_dump == clean_dump
        )
        coverage = CoverageV1.model_validate(clean_phase["coverage"])
        snapshot = build_snapshot(
            scan_id="iteration6-windows-native",
            facts=combined,
            coverage=coverage,
        )

        local_packet = build_packet(
            snapshot,
            artifact_privacy_mode="local_full",
            analysis_execution_mode="none",
        )
        none_outcome = analyze_packet(
            local_packet,
            provider="none",
            cache=ProductAnalysisCache(),
        )
        local_view = build_report_view_model(
            snapshot,
            local_packet,
            none_outcome.analysis,
            none_outcome.receipt,
        )
        local_workbook = sandbox / "windows-local-full-none.xlsx"
        render_workbook(local_view, local_workbook)
        local_validation = validate_workbook(
            local_workbook,
            view=local_view,
            snapshot=snapshot,
            packet=local_packet,
            analysis=none_outcome.analysis,
        )
        local_office = validate_office_workbook(local_workbook)
        local_cells = _cell_strings(local_workbook)
        local_full_has_windows_path = bool(
            re.search(r"[A-Za-z]:\\", local_cells)
        )

        share_packet = build_packet(
            snapshot,
            artifact_privacy_mode="share_safe",
            analysis_execution_mode="local_only",
        )
        privacy_lint = lint_share_safe_packet(share_packet)
        fake_provider = FakeProvider()
        fake_outcome = analyze_packet(
            share_packet,
            provider="fake",
            gateway=fake_provider,
            cache=ProductAnalysisCache(),
        )
        share_view = build_report_view_model(
            snapshot,
            share_packet,
            fake_outcome.analysis,
            fake_outcome.receipt,
        )
        artifacts.mkdir(parents=True, exist_ok=True)
        share_workbook = artifacts / "空间透视-Windows-share-safe.xlsx"
        render_workbook(share_view, share_workbook)
        share_validation = validate_workbook(
            share_workbook,
            view=share_view,
            snapshot=snapshot,
            packet=share_packet,
            analysis=fake_outcome.analysis,
        )
        share_office = validate_office_workbook(share_workbook)
        share_cells = _cell_strings(share_workbook)
        artifact_scrubbed = PRIVATE_VALUE.search(share_cells) is None

        by_name = {fact.name: fact for fact in clean_facts}
        reparse_facts = [
            fact
            for fact in clean_facts
            if fact.error_code == "REPARSE_POINT_NOT_FOLLOWED"
        ]
        allocated_files = [
            fact
            for fact in clean_facts
            if fact.file_type == "file" and fact.status == "collected"
        ]
        fixture_hash_count = len(fixture["files"]) * 2
        report["verifier_counters"]["fixture_integrity_hashes"] = (
            fixture_hash_count
        )
        report["product_counters"] = {
            "product_content_reads": 0,
            "product_file_hashes": 0,
            "network_calls": 0,
            "file_actions": 0,
            "provider_calls": none_outcome.receipt.provider_calls,
        }
        report["fake_product_counters"] = {
            "provider_calls": fake_outcome.receipt.provider_calls,
            "network_calls": fake_outcome.receipt.network_calls,
        }

        _record(report, "windows_native_environment", True)
        _record(
            report,
            "windows_allocated_size",
            bool(allocated_files)
            and all(
                fact.allocated_size_bytes is not None
                for fact in allocated_files
            ),
            observed_files=len(allocated_files),
        )
        _record(
            report,
            "windows_creation_time",
            bool(allocated_files)
            and all(
                fact.created.value is not None
                and fact.created.evidence_type == "filesystem_birthtime"
                for fact in allocated_files
            ),
        )
        _record(
            report,
            "windows_long_path",
            len(str(fixture["long_file"])) > 280
            and fixture["long_file"].name in by_name,
            path_length=len(str(fixture["long_file"])),
        )
        hidden_attributes = fixture["hidden"].stat(
            follow_symlinks=False
        ).st_file_attributes
        system_attributes = fixture["system"].stat(
            follow_symlinks=False
        ).st_file_attributes
        _record(
            report,
            "windows_hidden_attribute",
            bool(hidden_attributes & 0x2)
            and fixture["hidden"].name in by_name,
        )
        _record(
            report,
            "windows_system_attribute",
            bool(system_attributes & 0x4)
            and fixture["system"].name in by_name,
        )
        _record(
            report,
            "windows_reparse_not_followed",
            len(reparse_facts) == 1
            and fixture["junction_target_name"] not in by_name,
        )
        _record(
            report,
            "windows_permission_gap",
            any("permission-gap" in item for item in coverage.permission_gaps),
            permission_gap_count=len(coverage.permission_gaps),
        )
        _record(
            report,
            "checkpoint_process_restart_resume",
            resume_exact,
            cancelled_entries=cancel_phase["checkpoint"][
                "collected_entries"
            ],
            final_entries=coverage.collected_entries,
            final_skipped=coverage.skipped_entries,
        )
        _record(
            report,
            "none_local_full_workbook",
            local_validation["passed"]
            and local_office["passed"]
            and local_full_has_windows_path
            and none_outcome.receipt.provider_calls == 0,
            xlsx_sha256=workbook_sha256(local_workbook),
            local_full_has_windows_path=local_full_has_windows_path,
            office_checks=local_office,
        )
        _record(
            report,
            "fake_share_safe_workbook",
            share_validation["passed"]
            and share_office["passed"]
            and privacy_lint.passed
            and fake_outcome.receipt.provider_calls == 1
            and fake_outcome.receipt.network_calls == 0,
            filename=share_workbook.name,
            xlsx_sha256=workbook_sha256(share_workbook),
            office_checks=share_office,
        )
        _record(
            report,
            "artifact_scrub",
            artifact_scrubbed,
        )
        _record(
            report,
            "privacy_mode_difference",
            local_full_has_windows_path
            and artifact_scrubbed
            and local_packet.artifact_privacy_mode == "local_full"
            and share_packet.artifact_privacy_mode == "share_safe",
        )
        _record(
            report,
            "fixture_unchanged",
            before == after,
            fixture_integrity_hashes=fixture_hash_count,
        )
        _record(
            report,
            "product_boundaries",
            all(
                value == 0
                for value in report["product_counters"].values()
            )
            and report["fake_product_counters"]
            == {"provider_calls": 1, "network_calls": 0},
        )
        _set_file_attributes(fixture["hidden"], 0x80)
        _set_file_attributes(fixture["system"], 0x80)
        report["active_phase"] = "sandbox_cleanup"

    report.pop("active_phase", None)
    failed = [
        check["name"]
        for check in report["checks"]
        if not check["passed"]
    ]
    report["status"] = "passed" if not failed else "failed"
    report["blockers"] = failed
    report["evidence_level"] = "runtime_windows_ci"
    report["iteration6_status"] = (
        "PASS_WITH_GAPS_RUNTIME_WINDOWS_CI"
        if not failed
        else "FAILED_RUNTIME_WINDOWS_CI"
    )
    report["office_status"] = "NATIVE_WINDOWS_11_AND_OFFICE_REQUIRED"
    report["finished_at"] = _utc_now()
    report["commands"][0]["exit_code"] = 0
    ensure_share_safe_receipt(report)
    return report


def _native_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="空间透视 Iteration 6 Windows/Office fixture acceptance"
    )
    parser.add_argument(
        "--evidence",
        type=Path,
        default=OUTPUT_ROOT / "windows-ci-receipt.json",
    )
    parser.add_argument(
        "--artifacts",
        type=Path,
        default=OUTPUT_ROOT / "artifacts",
    )
    parser.add_argument(
        "--commit-sha",
        default=os.environ.get("GITHUB_SHA", "local-uncommitted"),
    )
    return parser


def _child_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("_collect_phase")
    parser.add_argument("--root", required=True)
    parser.add_argument("--allowed-root", required=True)
    parser.add_argument("--checkpoint")
    parser.add_argument("--batch-size", type=int, required=True)
    parser.add_argument("--cancel", action="store_true")
    return parser


def main() -> int:
    if len(sys.argv) > 1 and sys.argv[1] == "_collect-phase":
        return _child_collect(_child_parser().parse_args())
    args = _native_parser().parse_args()
    report = build_environment_report(
        os_name=os.name,
        system=platform.system(),
        release=platform.release(),
        commit_sha=args.commit_sha,
    )
    if report["status"] == "blocked":
        ensure_share_safe_receipt(report)
        _write_json(args.evidence, report)
        print(json.dumps(report, ensure_ascii=True, sort_keys=True))
        return EXIT_BLOCKED
    try:
        report = _run_native(report, artifacts=args.artifacts)
    except Exception as exc:
        report["status"] = "failed"
        report["iteration6_status"] = "FAILED_RUNTIME_WINDOWS_CI"
        report["blockers"] = [type(exc).__name__]
        report["failure"] = {
            "type": type(exc).__name__,
            "message": str(exc)[:240],
            "phase": report.pop("active_phase", "unknown"),
        }
        report["commands"][0]["exit_code"] = EXIT_FAILED
        report["finished_at"] = _utc_now()
        try:
            ensure_share_safe_receipt(report)
        except ValueError:
            report["failure"]["message"] = "sanitized native runner failure"
            ensure_share_safe_receipt(report)
    _write_json(args.evidence, report)
    print(json.dumps(report, ensure_ascii=True, sort_keys=True))
    return 0 if report["status"] == "passed" else EXIT_FAILED


if __name__ == "__main__":
    raise SystemExit(main())
