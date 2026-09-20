from __future__ import annotations

import builtins
from contextlib import ExitStack
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import socket
import stat
import sys
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from core.app.config import Settings
from core.app.main import create_app
from core.reports.models import SHEET_NAMES
from core.services.inventory import InventoryService


FIXTURE_ROOT = ROOT / "fixtures" / "ai-report"
STANDARD_MANIFEST = FIXTURE_ROOT / "standard" / "manifest.json"
TRIALS_ROOT = ROOT / "output" / "iteration4-synthetic-trials"
REQUIRED_FIXTURE_GUARD_SURFACES = {
    "builtins.open",
    "io.open",
    "Path.open",
    "Path.write_text",
    "Path.write_bytes",
    "Path.touch",
    "Path.mkdir",
    "Path.rmdir",
    "Path.chmod",
    "Path.symlink_to",
    "Path.hardlink_to",
    "Path.unlink",
    "Path.rename",
    "Path.replace",
    "os.open",
    "os.mkdir",
    "os.makedirs",
    "os.rmdir",
    "os.removedirs",
    "os.remove",
    "os.unlink",
    "os.rename",
    "os.replace",
    "os.symlink",
    "os.link",
    "os.utime",
    "os.truncate",
    "shutil.move",
    "shutil.copy",
    "shutil.copy2",
    "shutil.copyfile",
    "shutil.copytree",
}


def _consent_receipt_id(index: int) -> str:
    return f"consent_receipt_00000000-0000-4000-8000-{index:012d}"


def _client_action_id(index: int) -> str:
    return f"client_action_00000000-0000-4000-8000-{index:012d}"


def _validated_fixture_path(relative: str) -> Path:
    candidate = Path(relative)
    if candidate.is_absolute() or ".." in candidate.parts:
        raise RuntimeError(f"fixture manifest path is outside the allowlist: {relative}")
    current = FIXTURE_ROOT
    for part in candidate.parts:
        current = current / part
        info = current.lstat()
        if stat.S_ISLNK(info.st_mode):
            raise RuntimeError(f"fixture verifier refuses symlink: {relative}")
    return current


def _fixture_tree_metadata() -> list[dict[str, object]]:
    """Record the full fixture tree without following links or reading file content."""
    records: list[dict[str, object]] = []
    pending = [FIXTURE_ROOT]
    while pending:
        directory = pending.pop()
        directory_info = directory.lstat()
        records.append(
            {
                "relative_path": (
                    "." if directory == FIXTURE_ROOT else directory.relative_to(FIXTURE_ROOT).as_posix()
                ),
                "kind": "directory",
                "size": directory_info.st_size,
                "mtime_ns": directory_info.st_mtime_ns,
            }
        )
        with os.scandir(directory) as entries:
            ordered = sorted(entries, key=lambda item: item.name)
        for entry in ordered:
            path = Path(entry.path)
            info = entry.stat(follow_symlinks=False)
            if stat.S_ISLNK(info.st_mode):
                kind = "symlink"
            elif stat.S_ISDIR(info.st_mode):
                kind = "directory"
                pending.append(path)
                continue
            elif stat.S_ISREG(info.st_mode):
                kind = "file"
            else:
                kind = "other"
            records.append(
                {
                    "relative_path": path.relative_to(FIXTURE_ROOT).as_posix(),
                    "kind": kind,
                    "size": info.st_size,
                    "mtime_ns": info.st_mtime_ns,
                }
            )
    return sorted(records, key=lambda item: str(item["relative_path"]))


def _fixture_integrity() -> tuple[
    dict[str, list[dict[str, object]]],
    int,
    dict[str, object],
    dict[str, object],
]:
    """Verifier-only allowlist hashes plus content-free metadata for the full tree."""
    manifest_info = STANDARD_MANIFEST.lstat()
    if stat.S_ISLNK(manifest_info.st_mode) or not stat.S_ISREG(manifest_info.st_mode):
        raise RuntimeError("standard fixture manifest must be a regular file")
    manifest = json.loads(STANDARD_MANIFEST.read_text(encoding="utf-8"))
    entries = manifest.get("entries")
    if not isinstance(entries, list) or not entries or not all(
        isinstance(item, str) for item in entries
    ):
        raise RuntimeError("standard fixture manifest entries are invalid")
    allowed = ["standard/manifest.json", *entries]
    records: list[dict[str, object]] = []
    hashes = 0
    for relative in allowed:
        path = _validated_fixture_path(relative)
        info = path.lstat()
        if stat.S_ISREG(info.st_mode):
            kind = "file"
        elif stat.S_ISDIR(info.st_mode):
            kind = "directory"
        else:
            raise RuntimeError(f"fixture verifier refuses unsupported entry: {relative}")
        record: dict[str, object] = {
            "relative_path": relative,
            "kind": kind,
            "size": info.st_size,
            "mtime_ns": info.st_mtime_ns,
            "sha256": None,
        }
        if kind == "file":
            digest = hashlib.sha256()
            with path.open("rb") as source:
                while chunk := source.read(1024 * 1024):
                    digest.update(chunk)
            record["sha256"] = digest.hexdigest()
            hashes += 1
        records.append(record)
    tree_metadata = _fixture_tree_metadata()
    scope = {
        "fixture_id": manifest["fixture_id"],
        "fixture_version": manifest["fixture_version"],
        "hash_scope": "standard_manifest_allowlist_plus_manifest",
        "allowed_entries": len(allowed),
        "tree_entries": len(tree_metadata),
        "symlinks_followed": 0,
    }
    return {
        "allowlist_hashes": records,
        "tree_metadata": tree_metadata,
    }, hashes, scope, manifest


def _absolute_path(value: object) -> Path | None:
    if isinstance(value, int):
        return None
    try:
        return Path(os.path.abspath(os.fspath(value)))
    except TypeError:
        return None


def _is_fixture_payload(value: object) -> bool:
    path = _absolute_path(value)
    return path is not None and path.is_relative_to(FIXTURE_ROOT)


def _forbid_network(*_args: object, **_kwargs: object) -> None:
    raise AssertionError("network call is forbidden")


def _product_guards(
    stack: ExitStack,
    verified_manifest: dict[str, object],
) -> set[str]:
    original_builtin_open = builtins.open
    original_io_open = io.open
    original_path_open = Path.open
    original_path_write_text = Path.write_text
    original_path_write_bytes = Path.write_bytes
    original_path_touch = Path.touch
    original_path_mkdir = Path.mkdir
    original_path_rmdir = Path.rmdir
    original_path_chmod = Path.chmod
    original_path_symlink_to = Path.symlink_to
    original_path_hardlink_to = Path.hardlink_to
    original_path_unlink = Path.unlink
    original_path_rename = Path.rename
    original_path_replace = Path.replace
    original_os_open = os.open
    original_os_mkdir = os.mkdir
    original_os_makedirs = os.makedirs
    original_os_rmdir = os.rmdir
    original_os_removedirs = os.removedirs
    original_os_remove = os.remove
    original_os_unlink = os.unlink
    original_os_rename = os.rename
    original_os_replace = os.replace
    original_os_symlink = os.symlink
    original_os_link = os.link
    original_os_utime = os.utime
    original_os_truncate = os.truncate
    original_shutil_move = shutil.move
    original_shutil_copy = shutil.copy
    original_shutil_copy2 = shutil.copy2
    original_shutil_copyfile = shutil.copyfile
    original_shutil_copytree = shutil.copytree

    def guarded_builtin_open(file: object, *args: object, **kwargs: object):
        if _is_fixture_payload(file):
            raise AssertionError(f"product fixture open is forbidden: {file}")
        return original_builtin_open(file, *args, **kwargs)

    def guarded_io_open(file: object, *args: object, **kwargs: object):
        if _is_fixture_payload(file):
            raise AssertionError(f"product fixture open is forbidden: {file}")
        return original_io_open(file, *args, **kwargs)

    def guarded_path_open(self: Path, *args: object, **kwargs: object):
        if _is_fixture_payload(self):
            raise AssertionError(f"product fixture open is forbidden: {self}")
        return original_path_open(self, *args, **kwargs)

    def guard_mutation(*values: object) -> None:
        if any(_is_fixture_payload(value) for value in values):
            raise AssertionError(f"source fixture mutation is forbidden: {values}")

    def guarded_path_unlink(self: Path, *args: object, **kwargs: object):
        guard_mutation(self)
        return original_path_unlink(self, *args, **kwargs)

    def guarded_path_rename(self: Path, target: object):
        guard_mutation(self, target)
        return original_path_rename(self, target)

    def guarded_path_replace(self: Path, target: object):
        guard_mutation(self, target)
        return original_path_replace(self, target)

    def guarded_path_write_text(self: Path, *args: object, **kwargs: object):
        guard_mutation(self)
        return original_path_write_text(self, *args, **kwargs)

    def guarded_path_write_bytes(self: Path, *args: object, **kwargs: object):
        guard_mutation(self)
        return original_path_write_bytes(self, *args, **kwargs)

    def guarded_path_touch(self: Path, *args: object, **kwargs: object):
        guard_mutation(self)
        return original_path_touch(self, *args, **kwargs)

    def guarded_path_mkdir(self: Path, *args: object, **kwargs: object):
        guard_mutation(self)
        return original_path_mkdir(self, *args, **kwargs)

    def guarded_path_rmdir(self: Path, *args: object, **kwargs: object):
        guard_mutation(self)
        return original_path_rmdir(self, *args, **kwargs)

    def guarded_path_chmod(self: Path, *args: object, **kwargs: object):
        guard_mutation(self)
        return original_path_chmod(self, *args, **kwargs)

    def guarded_path_symlink_to(self: Path, target: object, *args: object, **kwargs: object):
        guard_mutation(self, target)
        return original_path_symlink_to(self, target, *args, **kwargs)

    def guarded_path_hardlink_to(self: Path, target: object):
        guard_mutation(self, target)
        return original_path_hardlink_to(self, target)

    def guarded_os_open(path: object, *args: object, **kwargs: object):
        if _is_fixture_payload(path):
            raise AssertionError(f"product fixture os.open is forbidden: {path}")
        return original_os_open(path, *args, **kwargs)

    def guarded_os_mkdir(path: object, *args: object, **kwargs: object):
        guard_mutation(path)
        return original_os_mkdir(path, *args, **kwargs)

    def guarded_os_makedirs(path: object, *args: object, **kwargs: object):
        guard_mutation(path)
        return original_os_makedirs(path, *args, **kwargs)

    def guarded_os_rmdir(path: object, *args: object, **kwargs: object):
        guard_mutation(path)
        return original_os_rmdir(path, *args, **kwargs)

    def guarded_os_removedirs(path: object, *args: object, **kwargs: object):
        guard_mutation(path)
        return original_os_removedirs(path, *args, **kwargs)

    def guarded_os_remove(path: object, *args: object, **kwargs: object):
        guard_mutation(path)
        return original_os_remove(path, *args, **kwargs)

    def guarded_os_unlink(path: object, *args: object, **kwargs: object):
        guard_mutation(path)
        return original_os_unlink(path, *args, **kwargs)

    def guarded_os_rename(src: object, dst: object, *args: object, **kwargs: object):
        guard_mutation(src, dst)
        return original_os_rename(src, dst, *args, **kwargs)

    def guarded_os_replace(src: object, dst: object, *args: object, **kwargs: object):
        guard_mutation(src, dst)
        return original_os_replace(src, dst, *args, **kwargs)

    def guarded_os_symlink(src: object, dst: object, *args: object, **kwargs: object):
        guard_mutation(src, dst)
        return original_os_symlink(src, dst, *args, **kwargs)

    def guarded_os_link(src: object, dst: object, *args: object, **kwargs: object):
        guard_mutation(src, dst)
        return original_os_link(src, dst, *args, **kwargs)

    def guarded_os_utime(path: object, *args: object, **kwargs: object):
        guard_mutation(path)
        return original_os_utime(path, *args, **kwargs)

    def guarded_os_truncate(path: object, *args: object, **kwargs: object):
        guard_mutation(path)
        return original_os_truncate(path, *args, **kwargs)

    def guarded_shutil_move(src: object, dst: object, *args: object, **kwargs: object):
        guard_mutation(src, dst)
        return original_shutil_move(src, dst, *args, **kwargs)

    def guarded_shutil_copy(src: object, dst: object, *args: object, **kwargs: object):
        guard_mutation(src, dst)
        return original_shutil_copy(src, dst, *args, **kwargs)

    def guarded_shutil_copy2(src: object, dst: object, *args: object, **kwargs: object):
        guard_mutation(src, dst)
        return original_shutil_copy2(src, dst, *args, **kwargs)

    def guarded_shutil_copyfile(src: object, dst: object, *args: object, **kwargs: object):
        guard_mutation(src, dst)
        return original_shutil_copyfile(src, dst, *args, **kwargs)

    def guarded_shutil_copytree(src: object, dst: object, *args: object, **kwargs: object):
        guard_mutation(src, dst)
        return original_shutil_copytree(src, dst, *args, **kwargs)

    stack.enter_context(patch.object(socket, "create_connection", side_effect=_forbid_network))
    stack.enter_context(patch.object(socket.socket, "connect", _forbid_network))
    stack.enter_context(patch.object(socket.socket, "connect_ex", _forbid_network))
    stack.enter_context(
        patch.object(
            InventoryService,
            "_load_manifest",
            side_effect=lambda: deepcopy(verified_manifest),
        )
    )
    stack.enter_context(patch.object(builtins, "open", guarded_builtin_open))
    stack.enter_context(patch.object(io, "open", guarded_io_open))
    stack.enter_context(patch.object(Path, "open", guarded_path_open))
    stack.enter_context(patch.object(Path, "write_text", guarded_path_write_text))
    stack.enter_context(patch.object(Path, "write_bytes", guarded_path_write_bytes))
    stack.enter_context(patch.object(Path, "touch", guarded_path_touch))
    stack.enter_context(patch.object(Path, "mkdir", guarded_path_mkdir))
    stack.enter_context(patch.object(Path, "rmdir", guarded_path_rmdir))
    stack.enter_context(patch.object(Path, "chmod", guarded_path_chmod))
    stack.enter_context(patch.object(Path, "symlink_to", guarded_path_symlink_to))
    stack.enter_context(patch.object(Path, "hardlink_to", guarded_path_hardlink_to))
    stack.enter_context(patch.object(Path, "unlink", guarded_path_unlink))
    stack.enter_context(patch.object(Path, "rename", guarded_path_rename))
    stack.enter_context(patch.object(Path, "replace", guarded_path_replace))
    stack.enter_context(patch.object(os, "open", guarded_os_open))
    stack.enter_context(patch.object(os, "mkdir", guarded_os_mkdir))
    stack.enter_context(patch.object(os, "makedirs", guarded_os_makedirs))
    stack.enter_context(patch.object(os, "rmdir", guarded_os_rmdir))
    stack.enter_context(patch.object(os, "removedirs", guarded_os_removedirs))
    stack.enter_context(patch.object(os, "remove", guarded_os_remove))
    stack.enter_context(patch.object(os, "unlink", guarded_os_unlink))
    stack.enter_context(patch.object(os, "rename", guarded_os_rename))
    stack.enter_context(patch.object(os, "replace", guarded_os_replace))
    stack.enter_context(patch.object(os, "symlink", guarded_os_symlink))
    stack.enter_context(patch.object(os, "link", guarded_os_link))
    stack.enter_context(patch.object(os, "utime", guarded_os_utime))
    stack.enter_context(patch.object(os, "truncate", guarded_os_truncate))
    stack.enter_context(patch.object(shutil, "move", guarded_shutil_move))
    stack.enter_context(patch.object(shutil, "copy", guarded_shutil_copy))
    stack.enter_context(patch.object(shutil, "copy2", guarded_shutil_copy2))
    stack.enter_context(patch.object(shutil, "copyfile", guarded_shutil_copyfile))
    stack.enter_context(patch.object(shutil, "copytree", guarded_shutil_copytree))
    return {
        "builtins.open",
        "io.open",
        "Path.open",
        "Path.write_text",
        "Path.write_bytes",
        "Path.touch",
        "Path.mkdir",
        "Path.rmdir",
        "Path.chmod",
        "Path.symlink_to",
        "Path.hardlink_to",
        "Path.unlink",
        "Path.rename",
        "Path.replace",
        "os.open",
        "os.mkdir",
        "os.makedirs",
        "os.rmdir",
        "os.removedirs",
        "os.remove",
        "os.unlink",
        "os.rename",
        "os.replace",
        "os.symlink",
        "os.link",
        "os.utime",
        "os.truncate",
        "shutil.move",
        "shutil.copy",
        "shutil.copy2",
        "shutil.copyfile",
        "shutil.copytree",
    }


def _post(
    client: TestClient,
    path: str,
    payload: dict[str, object],
) -> dict[str, object]:
    response = client.post(path, json=payload)
    if response.status_code not in {200, 201}:
        raise RuntimeError(f"{path} failed: {response.status_code} {response.text}")
    return response.json()


def _expect_post_error(
    client: TestClient,
    path: str,
    payload: dict[str, object],
    *,
    expected_status: int,
    expected_code: str,
) -> dict[str, object]:
    response = client.post(path, json=payload)
    if response.status_code != expected_status:
        raise RuntimeError(
            f"{path} expected {expected_status}, got {response.status_code}: {response.text}"
        )
    body = response.json()
    if body.get("code") != expected_code:
        raise RuntimeError(
            f"{path} expected {expected_code}, got {body.get('code')}: {response.text}"
        )
    return body


def _preview(
    client: TestClient,
    snapshot_id: str,
    *,
    top_k: int = 200,
    budget: dict[str, object] | None = None,
) -> dict[str, object]:
    payload: dict[str, object] = {
        "snapshot_id": snapshot_id,
        "artifact_privacy_mode": "share_safe",
        "analysis_execution_mode": "cloud_metadata_minimized",
        "top_k": top_k,
    }
    if budget is not None:
        payload["budget"] = budget
    preview = _post(
        client,
        "/v1/inventory/analysis-packets/preview",
        payload,
    )
    offer = preview.get("analysis_offer")
    if not isinstance(offer, dict):
        raise RuntimeError("synthetic preview did not return analysis_offer")
    return preview


def _consent(
    offer: dict[str, object],
    *,
    receipt_id: str,
    action_id: str,
) -> dict[str, object]:
    return {
        "consent_receipt_id": receipt_id,
        "client_action_id": action_id,
        "confirmed": True,
        "snapshot_id": offer["snapshot_id"],
        "snapshot_sha256": offer["snapshot_sha256"],
        "packet_id": offer["packet_id"],
        "packet_sha256": offer["packet_sha256"],
        "packet_bytes": offer["packet_bytes"],
        "artifact_privacy_mode": "share_safe",
        "analysis_execution_mode": "cloud_metadata_minimized",
        "provider": offer["provider"],
        "model": offer["model"],
        "prompt_version": offer["prompt_version"],
        "consent_schema_version": offer["consent_schema_version"],
        "analysis_schema_version": offer["analysis_schema_version"],
        "disclosure_version": offer["disclosure_version"],
        "budget": offer["budget"],
        "cost_basis": offer["cost_basis"],
    }


def _analysis_payload(
    offer: dict[str, object],
    *,
    receipt_id: str,
    action_id: str,
) -> dict[str, object]:
    return {
        "packet_id": offer["packet_id"],
        "provider": "fake",
        "consent": _consent(
            offer,
            receipt_id=receipt_id,
            action_id=action_id,
        ),
    }


def _report(
    client: TestClient,
    *,
    snapshot_id: str,
    packet_id: str,
    analysis_id: str,
) -> dict[str, object]:
    return _post(
        client,
        "/v1/inventory/reports",
        {
            "snapshot_id": snapshot_id,
            "packet_id": packet_id,
            "analysis_id": analysis_id,
        },
    )


def _verify_download(
    client: TestClient,
    report: dict[str, object],
) -> dict[str, object]:
    response = client.get(
        f"/v1/inventory/reports/{report['report_id']}/download"
    )
    if response.status_code != 200:
        raise RuntimeError(f"download failed: {response.status_code} {response.text}")
    expected_mime = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    if response.headers.get("content-type") != expected_mime:
        raise RuntimeError("download MIME does not match XLSX contract")
    if response.headers.get("cache-control") != "no-store":
        raise RuntimeError("download must use Cache-Control: no-store")
    digest = hashlib.sha256(response.content).hexdigest()
    manifest = report["manifest"]
    if digest != manifest["workbook_sha256"]:
        raise RuntimeError("download hash does not match report manifest")
    with ZipFile(BytesIO(response.content)) as archive:
        if archive.testzip() is not None:
            raise RuntimeError("download ZIP integrity check failed")
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    if workbook.sheetnames != SHEET_NAMES:
        raise RuntimeError("workbook sheets do not match frozen contract")
    workbook.close()
    if manifest["zip_integrity"] is not True:
        raise RuntimeError("report manifest did not assert ZIP integrity")
    if manifest["reopen_validated"] is not True:
        raise RuntimeError("report manifest did not assert workbook reopen")
    return {
        "report_id": report["report_id"],
        "analysis_status": report["status"],
        "workbook_sha256": digest,
        "mime": response.headers["content-type"],
        "cache_control": response.headers["cache-control"],
        "zip_integrity": True,
        "reopen_validated": True,
        "sheets": manifest["sheets"],
    }


def main() -> int:
    trial_id = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H-%M-%SZ-") + uuid4().hex[:12]
    output = TRIALS_ROOT / trial_id
    output.mkdir(parents=True, exist_ok=False)
    before, before_hashes, fixture_scope, verified_manifest = _fixture_integrity()
    settings = Settings(
        database_path=output / "iteration4.sqlite3",
        inventory_report_path=output / "reports",
    )
    with ExitStack() as product_stack:
        guarded_surfaces = _product_guards(product_stack, verified_manifest)
        with TestClient(create_app(settings)) as client:
            health = client.get("/health").json()
            capabilities = health["capabilities"]
            if capabilities["cloud_metadata_analysis_available"]:
                raise RuntimeError("real cloud capability must remain unavailable")
            if capabilities["cloud_metadata_analysis_enabled"]:
                raise RuntimeError("real cloud capability must remain disabled")
            if not capabilities["controlled_synthetic_consent_available"]:
                raise RuntimeError("controlled synthetic consent capability is unavailable")

            scan = _post(
                client,
                "/v1/inventory/scans",
                {
                    "fixture_id": "standard",
                    "profile": "inventory_metadata",
                    "compute_hash": False,
                },
            )
            snapshot = _post(
                client,
                "/v1/inventory/snapshots",
                {"scan_id": scan["scan_id"]},
            )

            success_preview = _preview(client, snapshot["snapshot_id"])
            success_offer = success_preview["analysis_offer"]
            if success_offer["network_calls"] != 0 or not success_offer["synthetic"]:
                raise RuntimeError("synthetic offer safety declaration is invalid")
            if success_offer["cost_basis"] != "synthetic_zero_external_cost":
                raise RuntimeError("synthetic cost basis is not explicit")
            if not success_preview["privacy_lint"]["passed"]:
                raise RuntimeError("share-safe synthetic packet failed privacy lint")

            calls_before_negative_inputs = client.app.state.inventory_service.fake_provider.calls
            missing_consent_error = _expect_post_error(
                client,
                "/v1/inventory/analyses",
                {"packet_id": success_offer["packet_id"], "provider": "fake"},
                expected_status=409,
                expected_code="ANALYSIS_CONSENT_REQUIRED",
            )
            rejected_payload = _analysis_payload(
                success_offer,
                receipt_id=_consent_receipt_id(1),
                action_id=_client_action_id(1),
            )
            rejected_payload["consent"]["confirmed"] = False
            rejected_consent_error = _expect_post_error(
                client,
                "/v1/inventory/analyses",
                rejected_payload,
                expected_status=422,
                expected_code="REQUEST_VALIDATION_FAILED",
            )
            unsafe_identifier_payload = _analysis_payload(
                success_offer,
                receipt_id="=1+1",
                action_id="@SUM(A1:A2)",
            )
            unsafe_identifier_error = _expect_post_error(
                client,
                "/v1/inventory/analyses",
                unsafe_identifier_payload,
                expected_status=422,
                expected_code="REQUEST_VALIDATION_FAILED",
            )
            if client.app.state.inventory_service.fake_provider.calls != calls_before_negative_inputs:
                raise RuntimeError("missing, rejected, or unsafe consent reached the provider")

            success_payload = _analysis_payload(
                success_offer,
                receipt_id=_consent_receipt_id(2),
                action_id=_client_action_id(2),
            )
            success = _post(client, "/v1/inventory/analyses", success_payload)
            calls_after_success = client.app.state.inventory_service.fake_provider.calls
            if success["receipt"]["provider_calls"] != 1:
                raise RuntimeError("synthetic success must call fake provider once")
            if success["receipt"]["network_calls"] != 0:
                raise RuntimeError("synthetic success must not call the network")
            if success["consent_receipt"]["status"] != "consumed":
                raise RuntimeError("synthetic consent was not consumed")

            idempotent = _post(client, "/v1/inventory/analyses", success_payload)
            calls_after_idempotent = client.app.state.inventory_service.fake_provider.calls
            if idempotent["analysis"]["analysis_id"] != success["analysis"]["analysis_id"]:
                raise RuntimeError("same client action did not return the original analysis")
            if calls_after_idempotent != calls_after_success:
                raise RuntimeError("idempotent retry called the provider again")

            changed_outer_preview = _preview(
                client,
                snapshot["snapshot_id"],
                top_k=4,
            )
            changed_outer_payload = deepcopy(success_payload)
            changed_outer_payload["packet_id"] = changed_outer_preview["analysis_offer"][
                "packet_id"
            ]
            changed_outer_error = _expect_post_error(
                client,
                "/v1/inventory/analyses",
                changed_outer_payload,
                expected_status=409,
                expected_code="ANALYSIS_CONSENT_REUSED",
            )
            if client.app.state.inventory_service.fake_provider.calls != calls_after_success:
                raise RuntimeError("changed outer packet replay reached the provider")

            cache_payload = _analysis_payload(
                success_offer,
                receipt_id=_consent_receipt_id(3),
                action_id=_client_action_id(3),
            )
            cached = _post(client, "/v1/inventory/analyses", cache_payload)
            if cached["receipt"]["status"] != "cache_hit":
                raise RuntimeError("new consent for unchanged packet did not hit product cache")
            if cached["receipt"]["provider_calls"] != 0:
                raise RuntimeError("product cache hit called the provider")
            if cached["receipt"]["cached_input_tokens"] != 0:
                raise RuntimeError("product cache hit must not claim provider cached tokens")

            mismatched = deepcopy(cache_payload)
            mismatched["consent"]["consent_receipt_id"] = _consent_receipt_id(4)
            mismatched["consent"]["client_action_id"] = _client_action_id(4)
            mismatched["consent"]["packet_sha256"] = "0" * 64
            calls_before_mismatch = client.app.state.inventory_service.fake_provider.calls
            mismatch_error = _expect_post_error(
                client,
                "/v1/inventory/analyses",
                mismatched,
                expected_status=409,
                expected_code="ANALYSIS_CONSENT_MISMATCH",
            )
            reused = deepcopy(success_payload)
            reused["consent"]["client_action_id"] = _client_action_id(5)
            reuse_error = _expect_post_error(
                client,
                "/v1/inventory/analyses",
                reused,
                expected_status=409,
                expected_code="ANALYSIS_CONSENT_REUSED",
            )
            if client.app.state.inventory_service.fake_provider.calls != calls_before_mismatch:
                raise RuntimeError("mismatched or reused consent reached the provider")

            budget = dict(success_offer["budget"])
            budget["max_input_tokens"] = 1
            budget_preview = _preview(
                client,
                snapshot["snapshot_id"],
                budget=budget,
            )
            budget_offer = budget_preview["analysis_offer"]
            budget_blocked = _post(
                client,
                "/v1/inventory/analyses",
                _analysis_payload(
                    budget_offer,
                    receipt_id=_consent_receipt_id(6),
                    action_id=_client_action_id(6),
                ),
            )
            if budget_blocked["receipt"]["status"] != "budget_blocked":
                raise RuntimeError("lower input budget did not block synthetic provider")
            if budget_blocked["receipt"]["provider_calls"] != 0:
                raise RuntimeError("budget blocked request called the provider")

            failure_preview = _preview(client, snapshot["snapshot_id"], top_k=1)
            failure_offer = failure_preview["analysis_offer"]
            client.app.state.inventory_service.fake_provider.fail = True
            failure = _post(
                client,
                "/v1/inventory/analyses",
                _analysis_payload(
                    failure_offer,
                    receipt_id=_consent_receipt_id(7),
                    action_id=_client_action_id(7),
                ),
            )
            client.app.state.inventory_service.fake_provider.fail = False

            schema_preview = _preview(client, snapshot["snapshot_id"], top_k=2)
            schema_offer = schema_preview["analysis_offer"]
            client.app.state.inventory_service.fake_provider.schema_failures = 2
            schema_fallback = _post(
                client,
                "/v1/inventory/analyses",
                _analysis_payload(
                    schema_offer,
                    receipt_id=_consent_receipt_id(8),
                    action_id=_client_action_id(8),
                ),
            )
            client.app.state.inventory_service.fake_provider.schema_failures = 0

            grounding_preview = _preview(client, snapshot["snapshot_id"], top_k=3)
            grounding_offer = grounding_preview["analysis_offer"]
            client.app.state.inventory_service.fake_provider.dangling_reference = True
            grounding_fallback = _post(
                client,
                "/v1/inventory/analyses",
                _analysis_payload(
                    grounding_offer,
                    receipt_id=_consent_receipt_id(9),
                    action_id=_client_action_id(9),
                ),
            )
            client.app.state.inventory_service.fake_provider.dangling_reference = False

            expected_fallbacks = (
                (failure, "provider_failed", "FAKE_PROVIDER_FAILED", 1),
                (schema_fallback, "schema_failed", "AI_SCHEMA_INVALID", 2),
                (grounding_fallback, "grounding_failed", "GROUNDING_FAILED", 1),
            )
            for outcome, status_value, error_code, provider_calls in expected_fallbacks:
                if outcome["analysis"]["status"] != "ai_failed_fallback":
                    raise RuntimeError(f"{status_value} did not produce fallback")
                if outcome["receipt"]["status"] != status_value:
                    raise RuntimeError(f"{status_value} receipt status is invalid")
                if outcome["receipt"]["error_code"] != error_code:
                    raise RuntimeError(f"{status_value} receipt error code is invalid")
                if outcome["receipt"]["provider_calls"] != provider_calls:
                    raise RuntimeError(f"{status_value} provider call count is invalid")
                if outcome["receipt"]["network_calls"] != 0:
                    raise RuntimeError(f"{status_value} used the network")
                if outcome["receipt"]["fallback"] is not True:
                    raise RuntimeError(f"{status_value} did not mark fallback")
                if outcome["consent_receipt"]["status"] != "consumed":
                    raise RuntimeError(f"{status_value} consent was not consumed")

            report_inputs = (
                (success_offer, success),
                (budget_offer, budget_blocked),
                (failure_offer, failure),
                (schema_offer, schema_fallback),
                (grounding_offer, grounding_fallback),
            )
            reports = []
            provider_calls_before_reports = client.app.state.inventory_service.fake_provider.calls
            success_report: dict[str, object] | None = None
            for offer, outcome in report_inputs:
                report = _report(
                    client,
                    snapshot_id=snapshot["snapshot_id"],
                    packet_id=offer["packet_id"],
                    analysis_id=outcome["analysis"]["analysis_id"],
                )
                if outcome is success:
                    success_report = report
                reports.append(_verify_download(client, report))
            if success_report is None:
                raise RuntimeError("success report was not created")
            repeated_report = _report(
                client,
                snapshot_id=snapshot["snapshot_id"],
                packet_id=success_offer["packet_id"],
                analysis_id=success["analysis"]["analysis_id"],
            )
            if repeated_report != success_report:
                raise RuntimeError("report rerender did not reuse the verified artifact")
            provider_calls_after_reports = client.app.state.inventory_service.fake_provider.calls
            if provider_calls_after_reports != provider_calls_before_reports:
                raise RuntimeError("report render called the provider")

        with TestClient(create_app(settings)) as restarted:
            restarted_analysis = restarted.get(
                f"/v1/inventory/analyses/{success['analysis']['analysis_id']}"
            )
            if restarted_analysis.status_code != 200:
                raise RuntimeError("synthetic analysis did not survive restart")
            if (
                restarted_analysis.json()["consent_receipt"]["consent_receipt_id"]
                != success["consent_receipt"]["consent_receipt_id"]
            ):
                raise RuntimeError("consent receipt did not survive restart")
            restarted_retry = _post(
                restarted,
                "/v1/inventory/analyses",
                success_payload,
            )
            if restarted_retry["analysis"]["analysis_id"] != success["analysis"]["analysis_id"]:
                raise RuntimeError("restart retry did not return original analysis")
            if restarted.app.state.inventory_service.fake_provider.calls != 0:
                raise RuntimeError("restart retry called the provider")

    after, after_hashes, after_scope, after_manifest = _fixture_integrity()
    analysis_results = [
        success,
        idempotent,
        cached,
        budget_blocked,
        failure,
        schema_fallback,
        grounding_fallback,
        restarted_retry,
    ]
    analysis_network_calls = sum(
        int(item["receipt"]["network_calls"]) for item in analysis_results
    )
    consent_network_calls = sum(
        int(item["consent_receipt"]["network_calls"])
        for item in analysis_results
        if item["consent_receipt"] is not None
    )
    hard_checks = {
        "fixture_unchanged": before == after,
        "fixture_tree_unchanged": (
            before["tree_metadata"] == after["tree_metadata"]
        ),
        "fixture_allowlist_hashes_unchanged": (
            before["allowlist_hashes"] == after["allowlist_hashes"]
        ),
        "fixture_scope_unchanged": fixture_scope == after_scope,
        "fixture_manifest_unchanged": verified_manifest == after_manifest,
        "fixture_hashes_recorded": before_hashes + after_hashes > 0,
        "fixture_guard_surface_complete": (
            guarded_surfaces == REQUIRED_FIXTURE_GUARD_SURFACES
        ),
        "product_content_reads_zero": scan["safety"]["product_content_reads"] == 0,
        "product_file_hashes_zero": scan["safety"]["product_file_hashes"] == 0,
        "product_network_calls_zero": scan["safety"]["network_calls"] == 0,
        "product_file_actions_zero": scan["safety"]["file_actions"] == 0,
        "analysis_network_calls_zero": analysis_network_calls == 0,
        "consent_network_calls_zero": consent_network_calls == 0,
        "idempotent_additional_provider_calls_zero": (
            calls_after_idempotent - calls_after_success == 0
        ),
        "cache_provider_calls_zero": cached["receipt"]["provider_calls"] == 0,
        "cache_provider_tokens_zero": cached["receipt"]["cached_input_tokens"] == 0,
        "budget_provider_calls_zero": budget_blocked["receipt"]["provider_calls"] == 0,
        "grounding_usage_preserved": (
            grounding_fallback["receipt"]["input_tokens"] > 0
            and grounding_fallback["receipt"]["output_tokens"] > 0
            and grounding_fallback["receipt"]["latency_ms"] > 0
            and grounding_fallback["consent_receipt"]["input_tokens"]
            == grounding_fallback["receipt"]["input_tokens"]
            and grounding_fallback["consent_receipt"]["output_tokens"]
            == grounding_fallback["receipt"]["output_tokens"]
            and grounding_fallback["consent_receipt"]["latency_ms"]
            == grounding_fallback["receipt"]["latency_ms"]
        ),
        "schema_retry_usage_preserved": (
            schema_fallback["receipt"]["provider_calls"] == 2
            and schema_fallback["receipt"]["input_tokens"] > 0
            and schema_fallback["receipt"]["output_tokens"] > 0
            and schema_fallback["receipt"]["latency_ms"] >= 2
            and schema_fallback["consent_receipt"]["input_tokens"]
            == schema_fallback["receipt"]["input_tokens"]
            and schema_fallback["consent_receipt"]["output_tokens"]
            == schema_fallback["receipt"]["output_tokens"]
            and schema_fallback["consent_receipt"]["latency_ms"]
            == schema_fallback["receipt"]["latency_ms"]
        ),
        "report_rerender_provider_calls_zero": (
            provider_calls_after_reports - provider_calls_before_reports == 0
        ),
        "real_cloud_capabilities_false": (
            capabilities["cloud_metadata_analysis_available"] is False
            and capabilities["cloud_metadata_analysis_enabled"] is False
        ),
        "all_xlsx_verified": all(
            item["zip_integrity"] and item["reopen_validated"] for item in reports
        ),
    }
    receipt = {
        "status": "passed" if all(hard_checks.values()) else "failed",
        "evidence_level": "fixture_e2e",
        "scope": "iteration_4a_controlled_synthetic_offline",
        "trial_directory": str(output),
        "scan_id": scan["scan_id"],
        "snapshot_id": snapshot["snapshot_id"],
        "fixture_hash_scope": fixture_scope,
        "hard_checks": hard_checks,
        "reports": reports,
        "packet": {
            "packet_id": success_offer["packet_id"],
            "packet_sha256": success_offer["packet_sha256"],
            "packet_bytes": success_offer["packet_bytes"],
            "privacy_lint": success_preview["privacy_lint"],
            "budget": success_offer["budget"],
            "cost_basis": success_offer["cost_basis"],
        },
        "consent": {
            "success_receipt_id": success["consent_receipt"]["consent_receipt_id"],
            "success_status": success["consent_receipt"]["status"],
            "cache_receipt_id": cached["consent_receipt"]["consent_receipt_id"],
            "cache_status": cached["consent_receipt"]["status"],
            "network_calls_total": consent_network_calls,
        },
        "model_calls": {
            "success": success["receipt"],
            "cache": cached["receipt"],
            "budget": budget_blocked["receipt"],
            "provider_fallback": failure["receipt"],
            "schema_fallback": schema_fallback["receipt"],
            "grounding_fallback": grounding_fallback["receipt"],
            "analysis_network_calls_total": analysis_network_calls,
            "idempotent_additional_provider_calls": (
                calls_after_idempotent - calls_after_success
            ),
            "report_rerender_additional_provider_calls": (
                provider_calls_after_reports - provider_calls_before_reports
            ),
        },
        "negative_gates": {
            "missing_consent": missing_consent_error["code"],
            "rejected_consent": rejected_consent_error["code"],
            "unsafe_identifier": unsafe_identifier_error["code"],
            "changed_outer_packet": changed_outer_error["code"],
            "mismatched_consent": mismatch_error["code"],
            "reused_consent": reuse_error["code"],
        },
        "product_content_reads": scan["safety"]["product_content_reads"],
        "product_file_hashes": scan["safety"]["product_file_hashes"],
        "network_calls": scan["safety"]["network_calls"],
        "file_actions": scan["safety"]["file_actions"],
        "fixture_integrity_hashes": before_hashes + after_hashes,
        "fixture_unchanged": before == after,
        "real_provider_spike": "BLOCKED_NOT_AUTHORIZED",
    }
    receipt_path = output / "synthetic-e2e-receipt.json"
    receipt_path.write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(receipt, ensure_ascii=False))
    return 0 if receipt["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
